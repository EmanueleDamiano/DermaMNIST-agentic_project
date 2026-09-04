"""DermaAgent workflow."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import random
import shutil
import sys
import tempfile
import urllib.request
import uuid
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterable

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import torch
from PIL import Image, ImageOps
from torch import nn
from torch.utils.data import DataLoader, Dataset


ROOT = Path(__file__).resolve().parent
PROJECT_ROOT = ROOT.parent
CONFIG_PATH = ROOT / "project_config.json"
DATA_PATH = ROOT / "data" / "dermamnist.npz"
REGISTRY_PATH = ROOT / "project_registry.jsonl"
MODEL_DIR = PROJECT_ROOT / "model"
MODEL_PATH = MODEL_DIR / "model.pt"
FINAL_OUTPUT = PROJECT_ROOT / "output" / "final_output"
FULL_PROCESS_OUTPUT = PROJECT_ROOT / "output" / "full_process"

WP_NAMES = {
    1: "Research and Requirements",
    2: "Data",
    3: "Training",
    4: "Evaluation",
    5: "Enrichment and XAI",
    6: "Testing",
    7: "Output",
    8: "Review",
}

DELIVERABLE_FILES = {
    "D1.1": "WP01/D1.1_research_review.md",
    "D1.2": "WP01/D1.2_requirements_matrix.csv",
    "D1.3": "WP01/D1.3_technical_baseline.json",
    "D2.1": "WP02/D2.1_dataset_registry.json",
    "D2.2": "WP02/D2.2_integrity_audit.json",
    "D2.3": "WP02/D2.3_preprocessing_pipeline.json",
    "D3.1": "WP03/D3.1_experiment_protocol.json",
    "D3.2": "WP03/D3.2_training_agent_specification.json",
    "D3.3": "WP03/D3.3_model_set_and_experiment_registry.json",
    "D4.1": "WP04/D4.1_evaluation_specification.json",
    "D4.2": "WP04/D4.2_evaluation_and_selection_report.json",
    "D4.3": "WP04/D4.3_frozen_model_manifest.json",
    "D5.1": "WP05/D5.1_source_policy_and_output_schema.json",
    "D5.2": "WP05/D5.2_enrichment_record.json",
    "D5.3": "WP05/D5.3_xai_record.json",
    "D6.1": "WP06/D6.1_testing_agent_specification.json",
    "D6.2": "WP06/D6.2_test_specification.json",
    "D6.3": "WP06/D6.3_final_test_report.json",
    "D7.1": "WP07/D7.1_explanation_agent_specification.json",
    "D7.2": "WP07/D7.2_case_report.json",
    "D7.3": "WP07/D7.3_prototype_package.json",
    "D8.1": "WP08/D8.1_review_checklist.json",
    "D8.2": "WP08/D8.2_acceptance_report.json",
}

ACTIVITIES = [
    ("LA1.1", "State of the art", "Reviewer", "D1.1", "All cited papers are linked to a requirement or design choice."),
    ("LA1.2", "Requirements", "Project Manager", "D1.2", "Every mandatory requirement has an owner, verification method, and deliverable."),
    ("LA1.3", "Feasibility baseline", "Technical Responsible", "D1.3", "Data, models, agents, metrics, and test rules are fixed before implementation."),
    ("LA2.1", "Dataset registry", "Technical Responsible", "D2.1", "10,015 images, seven labels, official partitions, dimensions, counts, and checksums verified."),
    ("LA2.2", "Integrity audit", "Reviewer", "D2.2", "Exact, near-duplicate, label, and cross-split findings logged with leakage policy."),
    ("LA2.3", "Data preparation", "Technical Responsible", "D2.3", "A clean run recreates split IDs, normalization, augmentation policy, and loaders."),
    ("LA3.1", "Model methods", "Technical Responsible", "D3.1", "Three classifier configurations use the same preprocessing, seeds, and metrics."),
    ("LA3.2", "Agent1", "Technical Responsible", "D3.2", "Agent1 trains the three candidates and records each run."),
    ("LA3.3", "Automated training", "Technical Responsible", "D3.3", "Every run is reproducible; failures retain explicit reasons."),
    ("LA4.1", "Metric policy", "Reviewer", "D4.1", "Metrics, tie-breaks, calibration, and uncertainty rules fixed before selection."),
    ("LA4.2", "Agent2 evaluation", "Technical Responsible", "D4.2", "Agent2 compares the three candidates and selects one model."),
    ("LA4.3", "Model selection and freeze", "Reviewer", "D4.3", "Validation-only selection stores checksums, versions, and selection record."),
    ("LA5.1", "Source and explanation policy", "Reviewer", "D5.1", "Approved sources, citations, unsupported-claim behavior, and schema defined."),
    ("LA5.2", "Agent3 class information", "Technical Responsible", "D5.2", "Agent3 adds cited class information without changing the prediction."),
    ("LA5.3", "XAI integration", "Technical Responsible", "D5.3", "Attribution comes from the frozen model and shares the prediction trace ID."),
    ("LA6.1", "Agent2 testing", "Technical Responsible", "D6.1", "Agent2 tests the selected model with the specified metrics."),
    ("LA6.2", "Test protocol", "Reviewer", "D6.2", "Test set, metrics, failure rules, and output format are fixed before execution."),
    ("LA6.3", "System test", "Reviewer", "D6.3", "Metrics, calibration, robustness, and failure cases recorded without tuning."),
    ("LA7.1", "Agent3 output", "Technical Responsible", "D7.1", "Agent3 combines prediction, attribution, and cited class information."),
    ("LA7.2", "Output format", "Project Manager", "D7.2", "Prediction, probabilities, uncertainty, attribution, citations, and trace ID present."),
    ("LA7.3", "Prototype run", "Project Manager", "D7.3", "An unseen benchmark case runs end to end without manual intermediate edits."),
    ("LA8.1", "Agent3 review", "Reviewer", "D8.1", "Agent3 checks the required records and acceptance results."),
    ("LA8.2", "Final assessment", "Reviewer", "D8.2", "No critical finding remains open; limitations include consequences."),
]

WORK_PACKAGES = [
    ("WP01", "Project goal", "Research, requirements, acceptance rules", "-", "Requirements approved"),
    ("WP02", "WP01 data requirements", "Dataset registry, audit, preprocessing", "WP05 source review", "Data package accepted"),
    ("WP03", "WP02 data package", "Model setup and Agent1 training", "WP05 source preparation", "Candidate models logged"),
    ("WP04", "WP03 model set", "Agent2 evaluation and model selection", "WP05 integration", "Selected model saved"),
    ("WP05", "WP01 source policy + WP03 output schema", "Agent3 class information and XAI", "WP03 and WP04", "Class information and XAI verified"),
    ("WP06", "WP04 selected model + WP05 module", "Agent2 final test", "-", "Test report accepted"),
    ("WP07", "WP06 accepted outputs", "Agent3 output and prototype run", "-", "Prototype demonstration passed"),
    ("WP08", "WP01-WP07 evidence", "Agent3 review and acceptance", "-", "Critical findings closed"),
]

AGENT_SPECS = {
    "Agent1": (
        "Train the three classifier candidates",
        ["train_candidate", "write_experiment_record"],
        "Cannot access final test labels",
    ),
    "Agent2": (
        "Evaluate candidates, select one model, and run the final test",
        ["evaluate_validation", "freeze_model", "verify_model_package", "run_final_test"],
        "Cannot change the model after selection",
    ),
    "Agent3": (
        "Create class information, attribution, output, and review records",
        ["retrieve_approved_fact", "compose_case_report", "review_evidence"],
        "Cannot change the prediction or probabilities",
    ),
}

ACCEPTANCE = [
    ("A1", "Dataset identity and split audit", "D2", "Official identity is verified and leakage findings have an explicit policy."),
    ("A2", "Agent roles and permissions", "D3-D8", "All logged tools are authorized and handoffs use versioned records."),
    ("A3", "Model selection", "D4", "The selection rule is set before testing and the selected model checksum is stored."),
    ("A4", "Class-balanced evaluation", "D4, D6", "Accuracy, macro, per-class, calibration, and confusion metrics exist."),
    ("A5", "Final test", "D6", "The final test runs after model selection with recorded settings and no retraining."),
    ("A6", "Explanation grounding", "D5, D7", "Prediction, attribution, cited facts, and trace ID are linked."),
    ("A7", "Prototype run", "D7", "An unseen benchmark sample runs end to end."),
    ("A8", "Final review", "D8", "No critical finding remains open and limitations are recorded."),
]

RISKS = [
    ("R1", "Cross-split duplicates or leakage", "High", "Duplicate clusters cross partitions", "Report official and leakage-aware results", "Reviewer"),
    ("R2", "Class imbalance hides weak classes", "High", "Aggregate and minority metrics diverge", "Select with macro metrics and report each class", "Technical + Reviewer"),
    ("R3", "Overfitting during repeated experiments", "High", "Training improves while validation stalls", "Fixed limits, early stopping, seeds, full history", "Technical"),
    ("R4", "Final test contamination", "Critical", "Test information appears before model selection", "Restrict test access and rerun any affected result", "Reviewer"),
    ("R5", "Invalid agent tool or parameters", "Medium", "Schema rejection in log", "Allow lists, typed parameters, explicit failure log", "Technical"),
    ("R6", "Wrong handoff record version", "High", "Trace or version mismatch", "Schema and trace checks block mismatched handoff", "Technical"),
    ("R7", "Unsupported explanation", "High", "Statement has no approved source", "Citation required; unsupported text omitted", "Reviewer"),
    ("R8", "XAI too coarse at 28×28", "Medium", "Diffuse or unstable attribution", "Label attribution coarse evidence only", "Technical"),
    ("R9", "Environment not reproducible", "Medium", "Clean run changes or dependencies fail", "Pinned dependencies, checksums, seeds, regression tests", "Technical"),
]

REFERENCES = [
    (1, "Yang et al.", "MedMNIST v2", "Scientific Data", 2023, "https://doi.org/10.1038/s41597-022-01721-8", "Dataset and benchmark baseline"),
    (2, "Abhishek, Jain, Hamarneh", "Investigating the Quality of DermaMNIST and Fitzpatrick17k", "Scientific Data", 2025, "https://doi.org/10.1038/s41597-025-04382-5", "Duplicate/leakage controls"),
    (3, "Yan et al.", "A multimodal vision foundation model for clinical dermatology", "Nature Medicine", 2025, "https://doi.org/10.1038/s41591-025-03747-y", "Scope distinction from large clinical models"),
    (4, "Chanda et al.", "Dermatologist-like explainable AI enhances melanoma diagnosis accuracy", "Nature Communications", 2025, "https://doi.org/10.1038/s41467-025-59532-5", "Explanation as evaluable output"),
    (5, "Dakhli and Barhoumi", "Improving skin lesion classification through saliency-guided loss functions", "Computers in Biology and Medicine", 2025, "https://doi.org/10.1016/j.compbiomed.2025.110299", "Classification and saliency"),
    (6, "Zuo, Wang, Wang", "Adaptive multimodal fusion for skin lesion classification", "Artificial Intelligence in Medicine", 2025, "https://doi.org/10.1016/j.artmed.2025.103091", "Uncertainty and per-class evaluation"),
    (7, "Nahm et al.", "Artificial Intelligence in Dermatology", "International Journal of Dermatology", 2025, "https://doi.org/10.1111/ijd.17847", "Clinical implementation limitations"),
    (8, "Han et al.", "Planet-wide performance of a skin disease AI algorithm", "npj Digital Medicine", 2025, "https://doi.org/10.1038/s41746-025-01980-w", "Distribution and real-world performance"),
    (9, "Li et al.", "Dataset nutrition label for dermatologic AI", "npj Digital Medicine", 2025, "https://doi.org/10.1038/s41746-025-02125-9", "Dataset documentation"),
    (10, "Ferber et al.", "Autonomous AI agent for oncology", "Nature Cancer", 2025, "https://doi.org/10.1038/s43018-025-00991-6", "Tool orchestration"),
    (11, "Chen et al.", "Multi-agent conversational diagnostic capability", "npj Digital Medicine", 2025, "https://doi.org/10.1038/s41746-025-01550-0", "Multi-agent reasoning"),
    (12, "Laiouar-Pedari et al.", "Prospective evidence on AI-assisted melanoma diagnostics", "JAMA Dermatology", 2026, "https://doi.org/10.1001/jamadermatol.2026.0217", "Prospective evidence limits"),
    (13, "Anriot et al.", "Limits of AI models for skin cancer diagnosis", "JAMA Dermatology", 2026, "https://doi.org/10.1001/jamadermatol.2026.1492", "Realistic-setting limitations"),
    (14, "Collaco et al.", "Agentic AI in healthcare: scoping review", "npj Digital Medicine", 2026, "https://doi.org/10.1038/s41746-026-02517-5", "Agent safety and validation gaps"),
    (15, "Yu et al.", "Multimodal AI agents in healthcare", "npj Digital Medicine", 2026, "https://doi.org/10.1038/s41746-026-03060-z", "Agent evidence maturity"),
    (16, "Schmidgall et al.", "AgentClinic", "npj Digital Medicine", 2026, "https://doi.org/10.1038/s41746-026-02674-7", "Sequential tool-use risk"),
]

REQUIREMENTS = [
    ("REQ-DATA-01", "Data", "Verify official file MD5, 10,015 records, 7 classes, 28×28 RGB, and official split sizes", "Technical Responsible", "Automated schema/count/checksum checks", "D2.1", "A1"),
    ("REQ-DATA-02", "Data", "Log exact duplicates, perceptual-hash candidates, label conflicts, and cross-split overlap", "Reviewer", "Full-dataset audit with affected counts/rates", "D2.2", "A1"),
    ("REQ-DATA-03", "Data", "Keep official benchmark split and define leakage-aware test exclusion without changing training", "Reviewer", "Policy plus parallel leakage-aware metrics", "D2.2", "A1"),
    ("REQ-DATA-04", "Data", "Compute normalization on training images only and augment training only", "Technical Responsible", "Pipeline configuration and code test", "D2.3", "A1"),
    ("REQ-TRAIN-01", "Training", "Compare three compact classifier families under one protocol", "Technical Responsible", "Three completed experiment records", "D3.1,D3.3", "A3"),
    ("REQ-TRAIN-02", "Training", "Record seeds, parameters, history, checkpoints, checksums, failures, and code version", "Technical Responsible", "Experiment registry completeness", "D3.3", "A2"),
    ("REQ-AGENT-01", "Agents", "Give Agent1, Agent2, and Agent3 explicit tasks, allowed tools, inputs, outputs, and limits", "Technical Responsible", "Specifications and unauthorized-tool unit test", "D3.2,D4.2,D5.2,D6.1,D7.1,D8.1", "A2"),
    ("REQ-TEST-01", "Testing", "Prevent Agent1 from receiving final test labels", "Reviewer", "Interface inspection and event sequence", "D3.2,D4.2,D6.3", "A5"),
    ("REQ-EVAL-01", "Evaluation", "Report accuracy, macro-F1, balanced accuracy, macro OvR AUROC, and per-class metrics", "Reviewer", "Metric schema checks", "D4.1,D4.2,D6.3", "A4"),
    ("REQ-EVAL-02", "Evaluation", "Report confusion matrix, ECE, Brier score, and NLL", "Reviewer", "Metric ranges, dimensions, and plots", "D4.1,D6.3", "A4"),
    ("REQ-EVAL-03", "Evaluation", "Select the highest validation macro-F1; use the listed tie-breakers when needed", "Agent2", "Selection record recomputation", "D4.2,D4.3", "A3"),
    ("REQ-EVAL-04", "Evaluation", "Set calibration and the uncertainty threshold from the validation split", "Agent2", "Selected-model manifest fields", "D4.2,D4.3", "A3,A5"),
    ("REQ-XAI-01", "XAI", "Generate attribution from the frozen classifier and link it to the prediction trace", "Technical Responsible", "Model checksum and trace equality", "D5.3", "A6"),
    ("REQ-SOURCE-01", "Enrichment", "Use approved sources with title and URL for every class fact", "Reviewer", "Citation schema validation", "D5.1,D5.2", "A6"),
    ("REQ-SOURCE-02", "Enrichment", "Never allow enrichment or explanation to modify class probabilities", "Technical Responsible", "Input/output equality assertion and unit test", "D5.2,D7.1,D7.2", "A2,A6"),
    ("REQ-OUTPUT-01", "Output", "Return prediction, probabilities, uncertainty, attribution, citations, and trace ID", "Project Manager", "Case-report schema check", "D7.2", "A6,A7"),
    ("REQ-ROBUST-01", "Testing", "Evaluate horizontal-flip and brightness perturbations", "Agent2", "Robustness metric records", "D6.2,D6.3", "A4"),
    ("REQ-REVIEW-01", "Review", "Block acceptance when any critical finding is open", "Agent3", "Acceptance decision logic test", "D8.1,D8.2", "A8"),
    ("REQ-REPRO-01", "Reproducibility", "Pin dependencies and record data, preprocessing, metric, model, and code versions", "Technical Responsible", "Clean verification plus version fields", "D1.3,D3.3,D4.3,D6.3", "A3,A5"),
    ("REQ-SCOPE-01", "Scope", "Label all outputs research-only; omit diagnosis and treatment advice", "Project Manager", "README, policy, and case-report text checks", "D5.1,D7.2", "A6"),
]


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def load_config() -> dict[str, Any]:
    with CONFIG_PATH.open(encoding="utf-8") as handle:
        return json.load(handle)


def wp_dir(number: int) -> Path:
    path = ROOT / "work_packages" / f"WP{number:02d}"
    path.mkdir(parents=True, exist_ok=True)
    return path


def deliverable_path(deliverable: str) -> Path:
    return ROOT / "work_packages" / DELIVERABLE_FILES[deliverable]


def saved_model_path(manifest: dict[str, Any]) -> Path:
    """Resolve the selected model stored in the top-level model folder."""
    path = (PROJECT_ROOT / manifest["model_file"]).resolve()
    if path != MODEL_PATH.resolve():
        raise ValueError("The trained model must be stored at model/model.pt")
    return path


def json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_safe(item) for item in value]
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return None if not np.isfinite(value) else float(value)
    if isinstance(value, Path):
        return value.as_posix()
    return value


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(json_safe(payload), indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    os.replace(temp, path)


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def file_hash(path: Path, algorithm: str = "sha256") -> str:
    digest = hashlib.new(algorithm)
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def data_md5(path: Path) -> str:
    return file_hash(path, "md5")


def set_seed(seed: int) -> None:
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(seed)
    torch.use_deterministic_algorithms(True, warn_only=True)


def append_event(agent: str, tool: str, status: str, details: dict[str, Any] | None = None) -> None:
    event = {"timestamp": now_iso(), "agent": agent, "tool": tool, "status": status, "details": details or {}}
    with REGISTRY_PATH.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(json_safe(event), ensure_ascii=False) + "\n")


class ControlledAgent:
    """Reject tools outside the agent allow list."""

    name = "Agent"
    allowed_tools: frozenset[str] = frozenset()

    def call(self, tool: str, function: Callable[..., Any], *args: Any, **kwargs: Any) -> Any:
        if tool not in self.allowed_tools:
            append_event(self.name, tool, "rejected", {"reason": "tool_not_allowed"})
            raise PermissionError(f"{self.name} is not allowed to call {tool}")
        append_event(self.name, tool, "started")
        try:
            result = function(*args, **kwargs)
        except Exception as exc:
            append_event(self.name, tool, "failed", {"error": f"{type(exc).__name__}: {exc}"})
            raise
        append_event(self.name, tool, "completed")
        return result


def download_dataset(config: dict[str, Any]) -> Path:
    DATA_PATH.parent.mkdir(parents=True, exist_ok=True)
    expected = config["data"]["md5"]
    if DATA_PATH.exists() and data_md5(DATA_PATH) == expected:
        return DATA_PATH
    if DATA_PATH.exists():
        DATA_PATH.unlink()
    request = urllib.request.Request(config["data"]["url"], headers={"User-Agent": "DermaAgent/1.0"})
    with tempfile.NamedTemporaryFile(delete=False, dir=DATA_PATH.parent, suffix=".part") as temp:
        temp_path = Path(temp.name)
        with urllib.request.urlopen(request, timeout=120) as response:
            shutil.copyfileobj(response, temp)
    if data_md5(temp_path) != expected:
        temp_path.unlink(missing_ok=True)
        raise ValueError("Downloaded dataset failed the official MD5 check")
    os.replace(temp_path, DATA_PATH)
    return DATA_PATH


def load_npz(path: Path = DATA_PATH) -> dict[str, np.ndarray]:
    with np.load(path, allow_pickle=False) as archive:
        return {key: archive[key] for key in archive.files}


def image_sha(image: np.ndarray) -> str:
    return hashlib.sha256(np.ascontiguousarray(image).tobytes()).hexdigest()


def perceptual_hash(image: np.ndarray) -> str:
    gray = Image.fromarray(image).convert("L").resize((9, 8), Image.Resampling.LANCZOS)
    reduced = np.asarray(gray, dtype=np.int16)
    bits = (reduced[:, 1:] >= reduced[:, :-1]).reshape(-1)
    value = sum(int(bit) << index for index, bit in enumerate(bits))
    return f"{value:016x}"


def split_checksum(images: np.ndarray, labels: np.ndarray) -> str:
    digest = hashlib.sha256()
    digest.update(np.ascontiguousarray(images).tobytes())
    digest.update(np.ascontiguousarray(labels).tobytes())
    return digest.hexdigest()


def audit_dataset(config: dict[str, Any], write_outputs: bool = True) -> tuple[dict[str, Any], dict[str, Any]]:
    path = download_dataset(config)
    arrays = load_npz(path)
    expected_keys = {f"{split}_{kind}" for split in ("train", "val", "test") for kind in ("images", "labels")}
    if set(arrays) != expected_keys:
        raise ValueError(f"Unexpected dataset keys: {sorted(arrays)}")

    records: list[dict[str, Any]] = []
    class_counts: dict[str, dict[str, int]] = {}
    split_checksums: dict[str, str] = {}
    shape_errors = 0
    label_errors = 0
    for split in ("train", "val", "test"):
        images = arrays[f"{split}_images"]
        labels = arrays[f"{split}_labels"].reshape(-1)
        if len(images) != config["data"]["expected_splits"][split]:
            raise ValueError(f"Unexpected {split} count: {len(images)}")
        shape_errors += int(sum(tuple(image.shape) != tuple(config["data"]["expected_shape"]) for image in images))
        label_errors += int(np.sum((labels < 0) | (labels >= 7)))
        class_counts[split] = {str(key): int(value) for key, value in sorted(Counter(labels.tolist()).items())}
        split_checksums[split] = split_checksum(images, labels)
        for index, (image, label) in enumerate(zip(images, labels, strict=True)):
            records.append({
                "sample_id": f"{split}-{index:05d}",
                "split": split,
                "index": index,
                "label": int(label),
                "sha256": image_sha(image),
                "perceptual_hash": perceptual_hash(image),
            })

    exact_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    perceptual_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        exact_groups[record["sha256"]].append(record)
        perceptual_groups[record["perceptual_hash"]].append(record)
    exact_clusters = [items for items in exact_groups.values() if len(items) > 1]
    near_clusters = [items for items in perceptual_groups.values() if len({item["sha256"] for item in items}) > 1]
    cross_exact = [items for items in exact_clusters if len({item["split"] for item in items}) > 1]
    cross_near = [items for items in near_clusters if len({item["split"] for item in items}) > 1]
    label_conflicts = [items for items in exact_clusters if len({item["label"] for item in items}) > 1]
    test_leakage = sorted({item["index"] for cluster in cross_exact for item in cluster if item["split"] == "test"})
    test_perceptual_overlap = sorted({item["index"] for cluster in cross_near for item in cluster if item["split"] == "test"})

    def summarize(clusters: list[list[dict[str, Any]]], limit: int = 100) -> dict[str, Any]:
        return {
            "cluster_count": len(clusters),
            "affected_records": int(sum(len(cluster) for cluster in clusters)),
            "shown_clusters": [[{key: item[key] for key in ("sample_id", "split", "label", "sha256", "perceptual_hash")} for item in cluster] for cluster in clusters[:limit]],
            "shown_cluster_limit": limit,
            "truncated": len(clusters) > limit,
        }

    total = sum(config["data"]["expected_splits"].values())
    registry = {
        "deliverable": "D2.1",
        "created_at": now_iso(),
        "dataset_id": config["data"]["dataset_id"],
        "source_url": config["data"]["url"],
        "license": config["data"]["license"],
        "file": DATA_PATH.relative_to(ROOT).as_posix(),
        "file_md5": data_md5(path),
        "file_sha256": file_hash(path),
        "sample_count": len(records),
        "expected_sample_count": total,
        "shape": config["data"]["expected_shape"],
        "split_counts": {split: len(arrays[f"{split}_images"]) for split in ("train", "val", "test")},
        "class_counts": class_counts,
        "split_checksums": split_checksums,
        "sample_id_policy": "{split}-{zero-padded index within official split}",
        "status": "verified" if len(records) == total and shape_errors == 0 and label_errors == 0 else "failed",
    }
    audit = {
        "deliverable": "D2.2",
        "created_at": now_iso(),
        "grain": "one record per official 28x28 DermaMNIST image",
        "intended_use": "benchmark training, model selection, and final testing",
        "checks": {
            "completeness": {"status": "pass" if len(records) == total else "fail", "actual": len(records), "expected": total},
            "shape_validity": {"status": "pass" if shape_errors == 0 else "fail", "invalid_records": shape_errors},
            "label_validity": {"status": "pass" if label_errors == 0 else "fail", "invalid_records": label_errors},
            "exact_duplicates": {"status": "finding" if exact_clusters else "pass", **summarize(exact_clusters)},
            "cross_split_exact_leakage": {"status": "finding" if cross_exact else "pass", "severity": "high" if cross_exact else "none", **summarize(cross_exact)},
            "perceptual_hash_candidates": {"status": "review", "note": "Identical 64-bit difference hashes are candidates, not proof of duplication.", **summarize(near_clusters)},
            "cross_split_perceptual_candidates": {"status": "review", **summarize(cross_near)},
            "exact_duplicate_label_conflicts": {"status": "finding" if label_conflicts else "pass", **summarize(label_conflicts)},
        },
        "test_exact_leakage_indices": test_leakage,
        "test_exact_leakage_rate": len(test_leakage) / config["data"]["expected_splits"]["test"],
        "test_perceptual_overlap_indices": test_perceptual_overlap,
        "test_perceptual_overlap_rate": len(test_perceptual_overlap) / config["data"]["expected_splits"]["test"],
        "policy": {
            "official": "Keep official splits unchanged for benchmark comparability.",
            "leakage_aware": "Also report test metrics after excluding exact test-image bytes present in train or validation.",
            "near_duplicates": "Retain difference-hash candidates because hash equality alone is not sufficient proof; report them for review.",
            "candidate_sensitivity": "Also report test metrics after excluding identical perceptual hashes across splits, labeled as sensitivity analysis rather than confirmed leakage.",
            "training_control": "No test image or label is used for training, calibration, uncertainty selection, or model choice.",
        },
        "impact": "Cross-split overlap can inflate benchmark estimates; parallel leakage-aware metrics make the impact visible.",
        "temporal_checks": "Not applicable: this static benchmark has no event or acquisition timestamp in the NPZ distribution.",
        "status": "accepted_with_documented_controls" if shape_errors == 0 and label_errors == 0 else "failed",
    }
    if write_outputs:
        write_json(deliverable_path("D2.1"), registry)
        write_json(deliverable_path("D2.2"), audit)
    return registry, audit


def stratified_indices(labels: np.ndarray, maximum: int | None, seed: int) -> np.ndarray:
    if maximum is None or maximum >= len(labels):
        return np.arange(len(labels))
    rng = np.random.default_rng(seed)
    chosen: list[int] = []
    counts = Counter(labels.reshape(-1).tolist())
    allocations = {label: max(1, int(maximum * count / len(labels))) for label, count in counts.items()}
    while sum(allocations.values()) > maximum:
        label = max(allocations, key=lambda item: allocations[item])
        if allocations[label] > 1:
            allocations[label] -= 1
    while sum(allocations.values()) < maximum:
        candidates = [label for label, count in counts.items() if allocations[label] < count]
        allocations[max(candidates, key=lambda item: counts[item] - allocations[item])] += 1
    for label, amount in allocations.items():
        pool = np.flatnonzero(labels.reshape(-1) == label)
        chosen.extend(rng.choice(pool, size=amount, replace=False).tolist())
    return np.array(sorted(chosen), dtype=np.int64)


@dataclass(frozen=True)
class TrainingBundle:
    train_images: np.ndarray
    train_labels: np.ndarray
    val_images: np.ndarray
    val_labels: np.ndarray
    train_indices: np.ndarray
    val_indices: np.ndarray


@dataclass(frozen=True)
class TestBundle:
    images: np.ndarray
    labels: np.ndarray
    indices: np.ndarray


class DatasetVault:
    """Only exposes labels through role-specific methods."""

    def __init__(self, config: dict[str, Any]):
        self.config = config
        self.path = download_dataset(config)

    def training_bundle(self) -> TrainingBundle:
        arrays = load_npz(self.path)
        train_idx = np.arange(len(arrays["train_images"]))
        val_idx = np.arange(len(arrays["val_images"]))
        return TrainingBundle(
            arrays["train_images"][train_idx], arrays["train_labels"][train_idx].reshape(-1),
            arrays["val_images"][val_idx], arrays["val_labels"][val_idx].reshape(-1), train_idx, val_idx,
        )

    def test_bundle(self, frozen_manifest: dict[str, Any]) -> TestBundle:
        model_path = saved_model_path(frozen_manifest)
        if file_hash(model_path) != frozen_manifest["model_sha256"]:
            raise ValueError("Frozen model checksum mismatch")
        if data_md5(self.path) != self.config["data"]["md5"]:
            raise ValueError("Dataset checksum mismatch")
        arrays = load_npz(self.path)
        indices = np.arange(len(arrays["test_images"]))
        return TestBundle(arrays["test_images"][indices], arrays["test_labels"][indices].reshape(-1), indices)

    def unlabeled_test_image(self, index: int) -> tuple[np.ndarray, str]:
        arrays = load_npz(self.path)
        if not 0 <= index < len(arrays["test_images"]):
            raise IndexError(f"Test index must be from 0 to {len(arrays['test_images']) - 1}")
        return arrays["test_images"][index], f"test-{index:05d}"


class ImageDataset(Dataset):
    def __init__(self, images: np.ndarray, labels: np.ndarray, mean: np.ndarray, std: np.ndarray, augment: bool = False):
        self.images = images
        self.labels = labels.astype(np.int64)
        self.mean = torch.tensor(mean, dtype=torch.float32).view(3, 1, 1)
        self.std = torch.tensor(std, dtype=torch.float32).view(3, 1, 1)
        self.augment = augment

    def __len__(self) -> int:
        return len(self.labels)

    def __getitem__(self, index: int) -> tuple[torch.Tensor, torch.Tensor]:
        image = torch.from_numpy(self.images[index]).permute(2, 0, 1).float() / 255.0
        if self.augment:
            if torch.rand(()) < 0.5:
                image = torch.flip(image, (2,))
            if torch.rand(()) < 0.2:
                image = torch.flip(image, (1,))
            image = torch.rot90(image, int(torch.randint(0, 4, ()).item()), (1, 2))
        return (image - self.mean) / self.std, torch.tensor(self.labels[index], dtype=torch.long)


class ConvNormAct(nn.Sequential):
    def __init__(self, in_channels: int, out_channels: int, kernel: int = 3, stride: int = 1, groups: int = 1):
        super().__init__(
            nn.Conv2d(in_channels, out_channels, kernel, stride, kernel // 2, groups=groups, bias=False),
            nn.BatchNorm2d(out_channels),
            nn.SiLU(inplace=True),
        )


class ResidualBlock(nn.Module):
    def __init__(self, in_channels: int, out_channels: int, stride: int = 1):
        super().__init__()
        self.body = nn.Sequential(
            nn.Conv2d(in_channels, out_channels, 3, stride, 1, bias=False), nn.BatchNorm2d(out_channels), nn.ReLU(inplace=True),
            nn.Conv2d(out_channels, out_channels, 3, 1, 1, bias=False), nn.BatchNorm2d(out_channels),
        )
        self.skip = nn.Identity() if stride == 1 and in_channels == out_channels else nn.Sequential(
            nn.Conv2d(in_channels, out_channels, 1, stride, bias=False), nn.BatchNorm2d(out_channels)
        )
        self.activation = nn.ReLU(inplace=True)

    def forward(self, x: torch.Tensor) -> torch.Tensor:
        return self.activation(self.body(x) + self.skip(x))


class ResNet18_28(nn.Module):
    def __init__(self, classes: int = 7, width: int = 12):
        super().__init__()
        self.stem = nn.Sequential(nn.Conv2d(3, width, 3, 1, 1, bias=False), nn.BatchNorm2d(width), nn.ReLU(inplace=True))
        layers: list[nn.Module] = []
        channels = width
        for out_channels, blocks, stride in ((width, 2, 1), (width * 2, 2, 2), (width * 4, 2, 2), (width * 8, 2, 2)):
            layers.append(ResidualBlock(channels, out_channels, stride))
            layers.extend(ResidualBlock(out_channels, out_channels) for _ in range(blocks - 1))
            channels = out_channels
        self.features = nn.Sequential(*layers)
        self.classifier = nn.Linear(channels, classes)

    def forward(self, x: torch.Tensor) -> torch.Tensor:
        x = self.features(self.stem(x)).mean(dim=(2, 3))
        return self.classifier(x)


class MBConv(nn.Module):
    def __init__(self, in_channels: int, out_channels: int, stride: int, expand: int = 3):
        super().__init__()
        hidden = in_channels * expand
        self.body = nn.Sequential(
            ConvNormAct(in_channels, hidden, 1), ConvNormAct(hidden, hidden, 3, stride, groups=hidden),
            nn.Conv2d(hidden, out_channels, 1, bias=False), nn.BatchNorm2d(out_channels),
        )
        self.residual = stride == 1 and in_channels == out_channels

    def forward(self, x: torch.Tensor) -> torch.Tensor:
        result = self.body(x)
        return result + x if self.residual else result


class EfficientNetB0_28(nn.Module):
    def __init__(self, classes: int = 7):
        super().__init__()
        self.features = nn.Sequential(
            ConvNormAct(3, 16), MBConv(16, 16, 1, 1), MBConv(16, 24, 2), MBConv(24, 24, 1),
            MBConv(24, 40, 2), MBConv(40, 40, 1), MBConv(40, 64, 2), ConvNormAct(64, 96, 1),
        )
        self.classifier = nn.Linear(96, classes)

    def forward(self, x: torch.Tensor) -> torch.Tensor:
        return self.classifier(self.features(x).mean(dim=(2, 3)))


class ConvNeXtBlock(nn.Module):
    def __init__(self, channels: int):
        super().__init__()
        self.block = nn.Sequential(
            nn.Conv2d(channels, channels, 7, padding=3, groups=channels), nn.GroupNorm(1, channels),
            nn.Conv2d(channels, channels * 3, 1), nn.GELU(), nn.Conv2d(channels * 3, channels, 1),
        )

    def forward(self, x: torch.Tensor) -> torch.Tensor:
        return x + self.block(x)


class ConvNeXtTiny_28(nn.Module):
    def __init__(self, classes: int = 7):
        super().__init__()
        self.features = nn.Sequential(
            nn.Conv2d(3, 20, 3, padding=1), ConvNeXtBlock(20),
            nn.Conv2d(20, 40, 2, stride=2), ConvNeXtBlock(40),
            nn.Conv2d(40, 80, 2, stride=2), ConvNeXtBlock(80), ConvNeXtBlock(80),
        )
        self.norm = nn.LayerNorm(80)
        self.classifier = nn.Linear(80, classes)

    def forward(self, x: torch.Tensor) -> torch.Tensor:
        x = self.features(x).mean(dim=(2, 3))
        return self.classifier(self.norm(x))


def make_model(name: str) -> nn.Module:
    models: dict[str, Callable[[], nn.Module]] = {
        "resnet18_28": ResNet18_28,
        "efficientnet_b0_28": EfficientNetB0_28,
        "convnext_tiny_28": ConvNeXtTiny_28,
    }
    if name not in models:
        raise ValueError(f"Unknown approved model: {name}")
    return models[name]()


def count_parameters(model: nn.Module) -> int:
    return sum(parameter.numel() for parameter in model.parameters() if parameter.requires_grad)


def softmax_numpy(logits: np.ndarray, temperature: float = 1.0) -> np.ndarray:
    scaled = logits / temperature
    scaled -= scaled.max(axis=1, keepdims=True)
    exp = np.exp(scaled)
    return exp / exp.sum(axis=1, keepdims=True)


def binary_auc(labels: np.ndarray, scores: np.ndarray) -> float:
    labels = labels.astype(bool)
    positives = int(labels.sum())
    negatives = len(labels) - positives
    if positives == 0 or negatives == 0:
        return float("nan")
    order = np.argsort(scores, kind="mergesort")
    ranks = np.empty(len(scores), dtype=float)
    sorted_scores = scores[order]
    start = 0
    while start < len(scores):
        end = start + 1
        while end < len(scores) and sorted_scores[end] == sorted_scores[start]:
            end += 1
        ranks[order[start:end]] = (start + 1 + end) / 2
        start = end
    return float((ranks[labels].sum() - positives * (positives + 1) / 2) / (positives * negatives))


def classification_metrics(labels: np.ndarray, probabilities: np.ndarray, bins: int = 10) -> dict[str, Any]:
    labels = labels.reshape(-1).astype(int)
    if probabilities.shape != (len(labels), 7):
        raise ValueError("Probabilities must have shape [N, 7]")
    predictions = probabilities.argmax(axis=1)
    confusion = np.zeros((7, 7), dtype=np.int64)
    np.add.at(confusion, (labels, predictions), 1)
    support = confusion.sum(axis=1)
    predicted_count = confusion.sum(axis=0)
    true_positive = np.diag(confusion).astype(float)
    precision = np.divide(true_positive, predicted_count, out=np.zeros(7), where=predicted_count > 0)
    recall = np.divide(true_positive, support, out=np.zeros(7), where=support > 0)
    f1 = np.divide(2 * precision * recall, precision + recall, out=np.zeros(7), where=(precision + recall) > 0)
    aucs = [binary_auc(labels == class_id, probabilities[:, class_id]) for class_id in range(7)]
    confidence = probabilities.max(axis=1)
    correct = predictions == labels
    ece = 0.0
    for left in np.linspace(0, 1, bins + 1)[:-1]:
        right = left + 1 / bins
        mask = (confidence >= left) & (confidence < right if right < 1 else confidence <= right)
        if mask.any():
            ece += mask.mean() * abs(float(correct[mask].mean()) - float(confidence[mask].mean()))
    one_hot = np.eye(7)[labels]
    clipped = np.clip(probabilities[np.arange(len(labels)), labels], 1e-12, 1)
    return {
        "sample_count": len(labels),
        "accuracy": float(correct.mean()),
        "macro_f1": float(f1.mean()),
        "balanced_accuracy": float(recall.mean()),
        "macro_auroc": float(np.nanmean(aucs)),
        "expected_calibration_error": float(ece),
        "brier_score": float(np.mean(np.sum((probabilities - one_hot) ** 2, axis=1))),
        "negative_log_likelihood": float(-np.log(clipped).mean()),
        "confusion_matrix": confusion.tolist(),
        "per_class": [{
            "class_id": index, "support": int(support[index]), "precision": float(precision[index]),
            "recall": float(recall[index]), "f1": float(f1[index]), "auroc": float(aucs[index]),
        } for index in range(7)],
    }


def attach_class_names(metrics: dict[str, Any], config: dict[str, Any]) -> dict[str, Any]:
    for item in metrics["per_class"]:
        item["class_name"] = config["data"]["labels"][str(item["class_id"])]
    return metrics


def fit_temperature(logits: np.ndarray, labels: np.ndarray) -> tuple[float, float]:
    best = (1.0, float("inf"))
    for temperature in np.geomspace(0.35, 4.0, 61):
        probabilities = softmax_numpy(logits, float(temperature))
        nll = float(-np.log(np.clip(probabilities[np.arange(len(labels)), labels.astype(int)], 1e-12, 1)).mean())
        if nll < best[1]:
            best = (float(temperature), nll)
    return best


def choose_uncertainty_threshold(labels: np.ndarray, probabilities: np.ndarray) -> dict[str, float]:
    predictions = probabilities.argmax(axis=1)
    errors = predictions != labels.reshape(-1)
    confidence = probabilities.max(axis=1)
    best_threshold, best_f1 = 0.5, -1.0
    for threshold in np.unique(np.quantile(confidence, np.linspace(0.05, 0.95, 37))):
        uncertain = confidence < threshold
        tp = int(np.sum(uncertain & errors))
        fp = int(np.sum(uncertain & ~errors))
        fn = int(np.sum(~uncertain & errors))
        score = 2 * tp / max(1, 2 * tp + fp + fn)
        if score > best_f1:
            best_threshold, best_f1 = float(threshold), float(score)
    return {"threshold": best_threshold, "validation_error_detection_f1": best_f1}


def predict_logits(model: nn.Module, images: np.ndarray, labels: np.ndarray, mean: np.ndarray, std: np.ndarray, batch_size: int, device: torch.device) -> tuple[np.ndarray, np.ndarray]:
    loader = DataLoader(ImageDataset(images, labels, mean, std), batch_size=batch_size, shuffle=False, num_workers=0)
    model.eval()
    logits: list[np.ndarray] = []
    truth: list[np.ndarray] = []
    with torch.no_grad():
        for inputs, targets in loader:
            logits.append(model(inputs.to(device)).cpu().numpy())
            truth.append(targets.numpy())
    return np.concatenate(logits), np.concatenate(truth)


def training_statistics(images: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    pixels = images.astype(np.float64) / 255.0
    mean = pixels.mean(axis=(0, 1, 2))
    std = np.maximum(pixels.std(axis=(0, 1, 2)), 1e-6)
    return mean, std


def train_candidate(
    name: str,
    bundle: TrainingBundle,
    config: dict[str, Any],
    target_accuracy: float,
    max_epochs: int,
    device: torch.device,
) -> dict[str, Any]:
    seed = config["training"]["seed"]
    set_seed(seed)
    mean, std = training_statistics(bundle.train_images)
    generator = torch.Generator().manual_seed(seed)
    train_loader = DataLoader(
        ImageDataset(bundle.train_images, bundle.train_labels, mean, std, augment=True),
        batch_size=config["training"]["batch_size"], shuffle=True, generator=generator, num_workers=0,
    )
    model = make_model(name).to(device)
    counts = np.bincount(bundle.train_labels, minlength=7)
    class_weights = len(bundle.train_labels) / (7 * np.maximum(counts, 1))
    criterion = nn.CrossEntropyLoss(weight=torch.tensor(class_weights, dtype=torch.float32, device=device))
    optimizer = torch.optim.AdamW(model.parameters(), lr=config["training"]["learning_rate"], weight_decay=config["training"]["weight_decay"])
    history: list[dict[str, Any]] = []
    best_score, best_state, stale = -1.0, None, 0
    stop_reason = "maximum epochs reached"
    for epoch in range(1, max_epochs + 1):
        model.train()
        total_loss = 0.0
        for inputs, targets in train_loader:
            inputs, targets = inputs.to(device), targets.to(device)
            optimizer.zero_grad(set_to_none=True)
            loss = criterion(model(inputs), targets)
            loss.backward()
            optimizer.step()
            total_loss += float(loss.item()) * len(targets)
        logits, truth = predict_logits(model, bundle.val_images, bundle.val_labels, mean, std, config["training"]["batch_size"], device)
        metrics = classification_metrics(truth, softmax_numpy(logits), config["evaluation"]["calibration_bins"])
        history.append({"epoch": epoch, "train_loss": total_loss / len(bundle.train_labels), "validation": metrics})
        if metrics["macro_f1"] > best_score + 1e-8:
            best_score = metrics["macro_f1"]
            best_state = {key: value.detach().cpu().clone() for key, value in model.state_dict().items()}
            stale = 0
        else:
            stale += 1
            if stale >= config["training"]["early_stopping_patience"]:
                stop_reason = "validation score stopped improving"
                break
        if metrics["accuracy"] >= target_accuracy:
            stop_reason = "target validation accuracy reached"
            break
    if best_state is None:
        raise RuntimeError("Training produced no checkpoint")
    checkpoint_dir = wp_dir(3) / "checkpoints"
    checkpoint_dir.mkdir(exist_ok=True)
    checkpoint_path = checkpoint_dir / f"{name}.pt"
    torch.save(best_state, checkpoint_path)
    return {
        "run_id": f"run-{name}-{seed}", "model_name": name, "status": "completed", "seed": seed,
        "target_validation_accuracy": target_accuracy, "maximum_epochs": max_epochs,
        "epochs_completed": len(history), "stop_reason": stop_reason,
        "train_sample_count": len(bundle.train_labels), "validation_sample_count": len(bundle.val_labels),
        "parameters": count_parameters(model), "checkpoint_file": checkpoint_path.relative_to(ROOT).as_posix(),
        "checkpoint_sha256": file_hash(checkpoint_path), "normalization_mean": mean.tolist(), "normalization_std": std.tolist(),
        "class_weights": class_weights.tolist(), "history": history, "best_validation_macro_f1": best_score,
        "created_at": now_iso(), "code_version": config["project"]["code_version"],
    }


class TrainingAgent(ControlledAgent):
    name = "Agent1"
    allowed_tools = frozenset({"train_candidate", "write_experiment_record"})

    def run(
        self,
        bundle: TrainingBundle,
        config: dict[str, Any],
        target_accuracy: float,
        max_epochs: int,
        device: torch.device,
        progress: Callable[[int, str], None] | None = None,
    ) -> list[dict[str, Any]]:
        records: list[dict[str, Any]] = []
        models = config["training"]["models"]
        for index, name in enumerate(models):
            if progress:
                progress(18 + index * 12, f"Agent1 is training {name} ({index + 1}/{len(models)})")
            try:
                records.append(self.call("train_candidate", train_candidate, name, bundle, config, target_accuracy, max_epochs, device))
            except Exception as exc:
                records.append({"run_id": f"run-{name}", "model_name": name, "status": "failed", "reason": str(exc), "created_at": now_iso()})
        self.call("write_experiment_record", write_json, deliverable_path("D3.3"), {
            "deliverable": "D3.3", "created_at": now_iso(),
            "target_validation_accuracy": target_accuracy, "maximum_epochs": max_epochs, "runs": records,
            "completed_count": sum(record["status"] == "completed" for record in records),
            "failed_count": sum(record["status"] == "failed" for record in records),
        })
        return records


def load_checkpoint_model(record: dict[str, Any], device: torch.device) -> nn.Module:
    if "model_file" in record:
        path = saved_model_path(record)
        expected_hash = record["model_sha256"]
    else:
        path = ROOT / record["checkpoint_file"]
        expected_hash = record["checkpoint_sha256"]
    if file_hash(path) != expected_hash:
        raise ValueError(f"Checkpoint checksum mismatch for {record['model_name']}")
    model = make_model(record["model_name"])
    model.load_state_dict(torch.load(path, map_location="cpu", weights_only=True))
    return model.to(device).eval()


def selection_key(item: dict[str, Any], config: dict[str, Any]) -> tuple[float, ...]:
    metrics = item["calibrated_metrics"]
    return tuple(float(metrics[key]) for key in [config["evaluation"]["primary_metric"], *config["evaluation"]["tie_breakers"]])


class EvaluationAgent(ControlledAgent):
    name = "Agent2"
    allowed_tools = frozenset({"evaluate_validation", "freeze_model"})

    def run(self, records: list[dict[str, Any]], bundle: TrainingBundle, config: dict[str, Any], device: torch.device) -> dict[str, Any]:
        completed = [record for record in records if record["status"] == "completed"]
        if len(completed) != len(config["training"]["models"]):
            raise RuntimeError("All three approved candidates must complete before selection")

        comparisons = []
        for record in completed:
            def evaluate() -> dict[str, Any]:
                model = load_checkpoint_model(record, device)
                logits, labels = predict_logits(
                    model, bundle.val_images, bundle.val_labels, np.array(record["normalization_mean"]),
                    np.array(record["normalization_std"]), config["training"]["batch_size"], device,
                )
                temperature, calibrated_nll = fit_temperature(logits, labels)
                raw = attach_class_names(classification_metrics(labels, softmax_numpy(logits), config["evaluation"]["calibration_bins"]), config)
                calibrated_probabilities = softmax_numpy(logits, temperature)
                calibrated = attach_class_names(classification_metrics(labels, calibrated_probabilities, config["evaluation"]["calibration_bins"]), config)
                uncertainty = choose_uncertainty_threshold(labels, calibrated_probabilities)
                return {
                    "run_id": record["run_id"], "model_name": record["model_name"], "checkpoint_sha256": record["checkpoint_sha256"],
                    "raw_metrics": raw, "temperature": temperature, "calibrated_validation_nll": calibrated_nll,
                    "calibrated_metrics": calibrated, "uncertainty": uncertainty,
                }
            comparisons.append(self.call("evaluate_validation", evaluate))

        selected = max(comparisons, key=lambda item: selection_key(item, config))
        class_prior = np.bincount(bundle.train_labels, minlength=7).astype(float)
        class_prior /= class_prior.sum()
        baseline_probabilities = np.repeat(class_prior[None, :], len(bundle.val_labels), axis=0)
        baseline_metrics = attach_class_names(classification_metrics(bundle.val_labels, baseline_probabilities, config["evaluation"]["calibration_bins"]), config)
        selection_report = {
            "deliverable": "D4.2", "created_at": now_iso(), "selection_data": "validation only",
            "policy": {"primary": config["evaluation"]["primary_metric"], "tie_breakers": config["evaluation"]["tie_breakers"]},
            "comparisons": comparisons, "selected_run_id": selected["run_id"], "selected_model": selected["model_name"],
            "selection_reason": "Lexicographic maximum of the predeclared validation metrics.",
            "project_baseline": {"method": "training_class_prior", "metrics": baseline_metrics},
            "competitive_on_primary_metric": selected["calibrated_metrics"][config["evaluation"]["primary_metric"]] > baseline_metrics[config["evaluation"]["primary_metric"]],
        }
        write_json(deliverable_path("D4.2"), selection_report)

        selected_record = next(record for record in completed if record["run_id"] == selected["run_id"])
        def freeze() -> dict[str, Any]:
            MODEL_DIR.mkdir(parents=True, exist_ok=True)
            shutil.copy2(ROOT / selected_record["checkpoint_file"], MODEL_PATH)
            manifest = {
                "deliverable": "D4.3", "frozen_at": now_iso(), "selection_report": deliverable_path("D4.2").relative_to(ROOT).as_posix(),
                "model_name": selected_record["model_name"], "run_id": selected_record["run_id"],
                "model_file": MODEL_PATH.relative_to(PROJECT_ROOT).as_posix(), "model_sha256": file_hash(MODEL_PATH),
                "source_checkpoint_sha256": selected_record["checkpoint_sha256"], "temperature": selected["temperature"],
                "uncertainty_threshold": selected["uncertainty"]["threshold"], "normalization_mean": selected_record["normalization_mean"],
                "normalization_std": selected_record["normalization_std"], "validation_metrics": selected["calibrated_metrics"],
                "dataset_id": config["data"]["dataset_id"], "preprocessing_id": "preprocess-v1",
                "metric_version": config["evaluation"]["metric_version"], "code_version": config["project"]["code_version"],
                "post_test_tuning_allowed": False,
            }
            write_json(deliverable_path("D4.3"), manifest)
            return manifest
        return self.call("freeze_model", freeze)


def perturb(images: np.ndarray, mode: str) -> np.ndarray:
    if mode == "horizontal_flip":
        return images[:, :, ::-1, :].copy()
    if mode == "brightness_minus_10_percent":
        return np.clip(images.astype(np.float32) * 0.9, 0, 255).astype(np.uint8)
    if mode == "brightness_plus_10_percent":
        return np.clip(images.astype(np.float32) * 1.1, 0, 255).astype(np.uint8)
    raise ValueError(f"Unknown fixed perturbation: {mode}")


def plot_confusion(confusion: list[list[int]], labels: dict[str, str], path: Path) -> None:
    matrix = np.array(confusion)
    figure, axis = plt.subplots(figsize=(8, 6))
    image = axis.imshow(matrix, cmap="Blues")
    figure.colorbar(image, ax=axis, fraction=0.046)
    short = [f"C{index}" for index in range(7)]
    axis.set(xticks=range(7), yticks=range(7), xticklabels=short, yticklabels=short, xlabel="Predicted", ylabel="True", title="DermaMNIST test confusion matrix")
    for row in range(7):
        for column in range(7):
            axis.text(column, row, str(matrix[row, column]), ha="center", va="center", color="white" if matrix[row, column] > matrix.max() / 2 else "black")
    figure.text(0.02, 0.01, " | ".join(f"C{k}: {v}" for k, v in labels.items()), fontsize=7)
    figure.tight_layout(rect=(0, 0.06, 1, 1))
    figure.savefig(path, dpi=160)
    plt.close(figure)


def plot_calibration(labels: np.ndarray, probabilities: np.ndarray, bins: int, path: Path) -> None:
    confidence = probabilities.max(axis=1)
    correct = probabilities.argmax(axis=1) == labels
    bin_accuracy, bin_confidence = [], []
    for left in np.linspace(0, 1, bins + 1)[:-1]:
        right = left + 1 / bins
        mask = (confidence >= left) & (confidence < right if right < 1 else confidence <= right)
        if mask.any():
            bin_accuracy.append(float(correct[mask].mean()))
            bin_confidence.append(float(confidence[mask].mean()))
    figure, axis = plt.subplots(figsize=(5, 5))
    axis.plot([0, 1], [0, 1], "--", color="gray", label="Perfect calibration")
    axis.plot(bin_confidence, bin_accuracy, "o-", color="#0b5cab", label="Selected model")
    axis.set(xlim=(0, 1), ylim=(0, 1), xlabel="Mean confidence", ylabel="Observed accuracy", title="Test reliability diagram")
    axis.legend()
    axis.grid(alpha=0.2)
    figure.tight_layout()
    figure.savefig(path, dpi=160)
    plt.close(figure)


class TestingAgent(ControlledAgent):
    name = "Agent2"
    allowed_tools = frozenset({"verify_model_package", "run_final_test"})

    def run(self, vault: DatasetVault, manifest: dict[str, Any], audit: dict[str, Any], config: dict[str, Any], device: torch.device) -> dict[str, Any]:
        self.call("verify_model_package", lambda: vault.test_bundle(manifest))

        def final_test() -> dict[str, Any]:
            bundle = vault.test_bundle(manifest)
            model_record = {
                "model_name": manifest["model_name"], "model_file": manifest["model_file"], "model_sha256": manifest["model_sha256"]
            }
            model = load_checkpoint_model(model_record, device)
            mean, std = np.array(manifest["normalization_mean"]), np.array(manifest["normalization_std"])
            logits, labels = predict_logits(model, bundle.images, bundle.labels, mean, std, config["training"]["batch_size"], device)
            probabilities = softmax_numpy(logits, manifest["temperature"])
            metrics = attach_class_names(classification_metrics(labels, probabilities, config["evaluation"]["calibration_bins"]), config)
            excluded_original = set(audit["test_exact_leakage_indices"])
            clean_mask = np.array([int(index) not in excluded_original for index in bundle.indices])
            leakage_metrics = attach_class_names(classification_metrics(labels[clean_mask], probabilities[clean_mask], config["evaluation"]["calibration_bins"]), config) if clean_mask.any() else None
            perceptual_overlap = set(audit["test_perceptual_overlap_indices"])
            sensitivity_mask = np.array([int(index) not in perceptual_overlap for index in bundle.indices])
            sensitivity_metrics = attach_class_names(classification_metrics(labels[sensitivity_mask], probabilities[sensitivity_mask], config["evaluation"]["calibration_bins"]), config) if sensitivity_mask.any() else None
            robustness = {}
            for mode in config["evaluation"]["robustness_checks"]:
                changed_logits, _ = predict_logits(model, perturb(bundle.images, mode), bundle.labels, mean, std, config["training"]["batch_size"], device)
                robustness[mode] = attach_class_names(classification_metrics(labels, softmax_numpy(changed_logits, manifest["temperature"]), config["evaluation"]["calibration_bins"]), config)
            confidence = probabilities.max(axis=1)
            predictions = probabilities.argmax(axis=1)
            failures = np.flatnonzero(predictions != labels)
            failures = failures[np.argsort(-confidence[failures])][: config["evaluation"]["failure_case_limit"]]
            failure_cases = [{
                "sample_id": f"test-{int(bundle.indices[index]):05d}", "true_class": int(labels[index]),
                "predicted_class": int(predictions[index]), "confidence": float(confidence[index]),
            } for index in failures]
            uncertainty = confidence < manifest["uncertainty_threshold"]
            test_dir = wp_dir(6)
            plot_confusion(metrics["confusion_matrix"], config["data"]["labels"], test_dir / "confusion_matrix.png")
            plot_calibration(labels, probabilities, config["evaluation"]["calibration_bins"], test_dir / "calibration_plot.png")
            report = {
                "deliverable": "D6.3", "created_at": now_iso(),
                "population": "Official DermaMNIST test split",
                "test_sample_count": len(labels), "test_indices_sha256": hashlib.sha256(bundle.indices.tobytes()).hexdigest(),
                "model_name": manifest["model_name"], "model_sha256": manifest["model_sha256"], "frozen_at": manifest["frozen_at"],
                "data_split_id": config["data"]["dataset_id"] + "-official-test", "preprocessing_id": manifest["preprocessing_id"],
                "metric_version": manifest["metric_version"], "code_version": manifest["code_version"],
                "temperature": manifest["temperature"], "uncertainty_threshold": manifest["uncertainty_threshold"],
                "uncertain_count": int(uncertainty.sum()), "uncertain_rate": float(uncertainty.mean()),
                "metrics": metrics, "leakage_aware": {
                    "excluded_exact_cross_split_test_records": int((~clean_mask).sum()),
                    "remaining_records": int(clean_mask.sum()), "metrics": leakage_metrics,
                },
                "perceptual_candidate_sensitivity": {
                    "interpretation": "Sensitivity analysis only; identical difference hashes are not confirmed duplicates.",
                    "excluded_candidate_test_records": int((~sensitivity_mask).sum()),
                    "remaining_records": int(sensitivity_mask.sum()), "metrics": sensitivity_metrics,
                },
                "robustness": robustness, "failure_cases": failure_cases,
                "plots": ["work_packages/WP06/confusion_matrix.png", "work_packages/WP06/calibration_plot.png"],
                "post_test_tuning_performed": False, "status": "completed_without_post_test_tuning",
            }
            write_json(deliverable_path("D6.3"), report)
            return report
        return self.call("run_final_test", final_test)


def xai_attribution(image: np.ndarray, manifest: dict[str, Any], device: torch.device, trace_id: str) -> dict[str, Any]:
    model_record = {"model_name": manifest["model_name"], "model_file": manifest["model_file"], "model_sha256": manifest["model_sha256"]}
    model = load_checkpoint_model(model_record, device)
    mean = torch.tensor(manifest["normalization_mean"], dtype=torch.float32).view(1, 3, 1, 1)
    std = torch.tensor(manifest["normalization_std"], dtype=torch.float32).view(1, 3, 1, 1)
    tensor = torch.from_numpy(image).permute(2, 0, 1).unsqueeze(0).float() / 255.0
    tensor = ((tensor - mean) / std).to(device).requires_grad_(True)
    logits = model(tensor)
    predicted = int(logits.argmax(dim=1).item())
    model.zero_grad(set_to_none=True)
    logits[0, predicted].backward()
    attribution = tensor.grad.detach().abs().mean(dim=1)[0].cpu().numpy()
    attribution -= attribution.min()
    attribution /= max(float(attribution.max()), 1e-12)
    npy_path = wp_dir(5) / "attribution.npy"
    np.save(npy_path, attribution)
    overlay_path = wp_dir(5) / "attribution_overlay.png"
    figure, axis = plt.subplots(figsize=(4, 4))
    axis.imshow(image)
    axis.imshow(attribution, cmap="inferno", alpha=0.55, vmin=0, vmax=1)
    axis.set_title("Input-gradient attribution")
    axis.axis("off")
    figure.tight_layout()
    figure.savefig(overlay_path, dpi=180)
    plt.close(figure)
    return {
        "deliverable": "D5.3", "created_at": now_iso(), "trace_id": trace_id, "method": "absolute_input_gradient",
        "target_class": predicted, "model_sha256": manifest["model_sha256"], "array_file": npy_path.relative_to(ROOT).as_posix(),
        "overlay_file": overlay_path.relative_to(ROOT).as_posix(), "minimum": float(attribution.min()), "maximum": float(attribution.max()),
        "interpretation": "Coarse evidence of input sensitivity at 28×28; not a lesion boundary or clinical localization.",
    }


def model_prediction(image: np.ndarray, sample_id: str, manifest: dict[str, Any], config: dict[str, Any], device: torch.device, trace_id: str) -> dict[str, Any]:
    record = {"model_name": manifest["model_name"], "model_file": manifest["model_file"], "model_sha256": manifest["model_sha256"]}
    model = load_checkpoint_model(record, device)
    logits, _ = predict_logits(model, image[None], np.array([0]), np.array(manifest["normalization_mean"]), np.array(manifest["normalization_std"]), 1, device)
    probabilities = softmax_numpy(logits, manifest["temperature"])[0]
    predicted = int(probabilities.argmax())
    return {
        "trace_id": trace_id, "sample_id": sample_id, "model_name": manifest["model_name"], "model_sha256": manifest["model_sha256"],
        "predicted_class_id": predicted, "predicted_class": config["data"]["labels"][str(predicted)],
        "probabilities": {config["data"]["labels"][str(index)]: float(value) for index, value in enumerate(probabilities)},
        "confidence": float(probabilities[predicted]), "uncertainty_threshold": manifest["uncertainty_threshold"],
        "uncertain": bool(probabilities[predicted] < manifest["uncertainty_threshold"]),
    }


class EnrichmentAgent(ControlledAgent):
    name = "Agent3"
    allowed_tools = frozenset({"retrieve_approved_fact"})

    def run(self, prediction: dict[str, Any], config: dict[str, Any]) -> dict[str, Any]:
        original = json.dumps(prediction, sort_keys=True)
        class_id = str(prediction["predicted_class_id"])

        def retrieve() -> dict[str, Any]:
            source = config["sources"][class_id]
            return {
                "deliverable": "D5.2", "created_at": now_iso(), "trace_id": prediction["trace_id"],
                "predicted_class_id": int(class_id), "fact": source["fact"], "citations": source["citations"],
                "classifier_mutation": False, "source_policy": deliverable_path("D5.1").relative_to(ROOT).as_posix(),
            }
        result = self.call("retrieve_approved_fact", retrieve)
        if json.dumps(prediction, sort_keys=True) != original:
            raise RuntimeError("Enrichment mutated the classifier prediction")
        write_json(deliverable_path("D5.2"), result)
        return result


class ExplanationAgent(ControlledAgent):
    name = "Agent3"
    allowed_tools = frozenset({"compose_case_report"})

    def run(self, prediction: dict[str, Any], xai: dict[str, Any], enrichment: dict[str, Any]) -> dict[str, Any]:
        if len({prediction["trace_id"], xai["trace_id"], enrichment["trace_id"]}) != 1:
            raise ValueError("Trace mismatch at explanation handoff")
        original_probabilities = dict(prediction["probabilities"])

        def compose() -> dict[str, Any]:
            return {
                "deliverable": "D7.2", "created_at": now_iso(), "trace_id": prediction["trace_id"], "sample_id": prediction["sample_id"],
                "research_only_notice": "Research benchmark output only; not a diagnosis or treatment recommendation.",
                "model_evidence": prediction, "xai_evidence": xai,
                "cited_class_information": {"fact": enrichment["fact"], "citations": enrichment["citations"]},
                "explanation": (
                    f"Predicted class: {prediction['predicted_class']}. Confidence: {prediction['confidence']:.3f}. "
                    f"Uncertainty: {'flagged' if prediction['uncertain'] else 'not flagged'}. "
                    f"The attribution image shows input sensitivity. Class information: {enrichment['fact']}"
                ),
                "unsupported_statements_omitted": True,
            }
        report = self.call("compose_case_report", compose)
        if report["model_evidence"]["probabilities"] != original_probabilities:
            raise RuntimeError("Explanation altered class probabilities")
        write_json(deliverable_path("D7.2"), report)
        return report


def report_progress(callback: Callable[[int, str], None] | None, percent: int, stage: str) -> None:
    if callback:
        callback(percent, stage)


def build_case(
    image: np.ndarray,
    sample_id: str,
    manifest: dict[str, Any],
    config: dict[str, Any],
    device: torch.device,
    progress: Callable[[int, str], None] | None = None,
) -> dict[str, Any]:
    trace_id = str(uuid.uuid4())
    report_progress(progress, 20, "Agent2 is classifying the image")
    prediction = model_prediction(image, sample_id, manifest, config, device, trace_id)
    report_progress(progress, 45, "Agent3 is creating the attribution")
    xai = xai_attribution(image, manifest, device, trace_id)
    write_json(deliverable_path("D5.3"), xai)
    report_progress(progress, 70, "Agent3 is adding class information")
    enrichment = EnrichmentAgent().run(prediction, config)
    report_progress(progress, 85, "Agent3 is building the result")
    report = ExplanationAgent().run(prediction, xai, enrichment)
    package = {
        "deliverable": "D7.3", "created_at": now_iso(), "trace_id": trace_id, "sample_id": sample_id, "status": "passed",
        "manual_intermediate_editing": False, "model_sha256": manifest["model_sha256"],
        "outputs": {
            "case_report": deliverable_path("D7.2").relative_to(ROOT).as_posix(),
            "attribution": xai["overlay_file"], "enrichment": deliverable_path("D5.2").relative_to(ROOT).as_posix(),
        },
    }
    package["output_sha256"] = {name: file_hash(ROOT / path) for name, path in package["outputs"].items()}
    write_json(deliverable_path("D7.3"), package)
    return report


def run_case(vault: DatasetVault, manifest: dict[str, Any], config: dict[str, Any], device: torch.device, index: int = 0) -> dict[str, Any]:
    image, sample_id = vault.unlabeled_test_image(index)
    return build_case(image, sample_id, manifest, config, device)


def validate_input_image(path: Path) -> None:
    if not path.is_file():
        raise FileNotFoundError(f"Image not found: {path}")
    if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".bmp", ".webp"}:
        raise ValueError("Use a JPG, JPEG, PNG, BMP, or WebP image")
    try:
        with Image.open(path) as source:
            source.verify()
    except OSError as exc:
        raise ValueError(f"Cannot read image: {path}") from exc


def run_user_image(
    path: Path,
    manifest: dict[str, Any],
    config: dict[str, Any],
    device: torch.device,
    output_dir: Path,
    progress: Callable[[int, str], None] | None = None,
) -> dict[str, Any]:
    report_progress(progress, 5, "Reading the input image")
    validate_input_image(path)
    try:
        with Image.open(path) as source:
            image = np.asarray(ImageOps.exif_transpose(source).convert("RGB").resize((28, 28), Image.Resampling.LANCZOS), dtype=np.uint8).copy()
    except OSError as exc:
        raise ValueError(f"Cannot read image: {path}") from exc
    report_progress(progress, 10, "Resizing the image to 28 x 28 RGB")
    sample_id = "input-" + file_hash(path)[:12]
    report = build_case(image, sample_id, manifest, config, device, progress)
    output_dir.mkdir(parents=True, exist_ok=True)
    attribution_path = output_dir / "attribution.png"
    shutil.copy2(ROOT / report["xai_evidence"]["overlay_file"], attribution_path)
    result = {
        "trace_id": report["trace_id"],
        "input": str(path.resolve()),
        "prediction": report["model_evidence"]["predicted_class"],
        "confidence": report["model_evidence"]["confidence"],
        "uncertain": report["model_evidence"]["uncertain"],
        "probabilities": report["model_evidence"]["probabilities"],
        "class_information": report["cited_class_information"]["fact"],
        "citations": report["cited_class_information"]["citations"],
        "attribution": "attribution.png",
        "notice": "Research and education only.",
    }
    write_json(output_dir / "prediction.json", result)
    report_progress(progress, 100, "Result ready")
    return result


def load_events() -> list[dict[str, Any]]:
    if not REGISTRY_PATH.exists():
        return []
    return [json.loads(line) for line in REGISTRY_PATH.read_text(encoding="utf-8").splitlines() if line.strip()]


class ReviewAgent(ControlledAgent):
    name = "Agent3"
    allowed_tools = frozenset({"review_evidence"})

    def run(self, config: dict[str, Any]) -> dict[str, Any]:
        def review() -> tuple[dict[str, Any], dict[str, Any]]:
            findings: list[dict[str, Any]] = []
            missing = [deliverable for deliverable in DELIVERABLE_FILES if deliverable not in {"D8.1", "D8.2"} and not deliverable_path(deliverable).exists()]
            if missing:
                findings.append({"severity": "critical", "issue": "Missing required deliverables", "evidence": missing})
            audit = read_json(deliverable_path("D2.2"))
            selection = read_json(deliverable_path("D4.2"))
            manifest = read_json(deliverable_path("D4.3"))
            test = read_json(deliverable_path("D6.3"))
            case = read_json(deliverable_path("D7.2"))
            prototype = read_json(deliverable_path("D7.3"))
            events = load_events()
            unauthorized = [event for event in events if event["status"] == "completed" and event["tool"] not in AGENT_SPECS.get(event["agent"], (None, [], None))[1]]
            metric_keys = {"accuracy", "macro_f1", "balanced_accuracy", "macro_auroc", "expected_calibration_error", "brier_score", "negative_log_likelihood", "confusion_matrix", "per_class"}
            trace_linked = case["trace_id"] == case["model_evidence"]["trace_id"] == case["xai_evidence"]["trace_id"] == prototype["trace_id"]
            citations_ok = bool(case["cited_class_information"]["citations"]) and all(item.get("url") and item.get("title") for item in case["cited_class_information"]["citations"])
            checks = [
                ("A1", read_json(deliverable_path("D2.1"))["status"] == "verified" and audit["status"] == "accepted_with_documented_controls", "Dataset identity, integrity findings, and parallel leakage-aware policy recorded."),
                ("A2", not unauthorized and all(agent in {event["agent"] for event in events} for agent in AGENT_SPECS), "Agent1, Agent2, and Agent3 logged only allowed actions."),
                ("A3", selection["selection_data"] == "validation only" and file_hash(saved_model_path(manifest)) == manifest["model_sha256"] and manifest["post_test_tuning_allowed"] is False, "Selection rule and model checksum verified."),
                ("A4", metric_keys.issubset(test["metrics"]) and len(test["metrics"]["per_class"]) == 7 and len(test["metrics"]["confusion_matrix"]) == 7, "Aggregate, class-balanced, per-class, calibration, and confusion evidence present."),
                ("A5", test["post_test_tuning_performed"] is False and test["frozen_at"] == manifest["frozen_at"], "Final test used the saved model and no retraining is recorded."),
                ("A6", trace_linked and citations_ok and case["unsupported_statements_omitted"], "Prediction, XAI, and approved sources share one trace."),
                ("A7", prototype["status"] == "passed" and all((ROOT / path).exists() for path in prototype["outputs"].values()), "Unseen benchmark case produced all declared outputs."),
            ]
            for identifier, passed, evidence in checks:
                if not passed:
                    findings.append({"severity": "critical", "issue": f"{identifier} failed", "evidence": evidence})
            critical_open = sum(item["severity"] == "critical" for item in findings)
            checks.append(("A8", critical_open == 0, "No critical findings remain open." if critical_open == 0 else f"{critical_open} critical findings remain."))
            checklist = {
                "deliverable": "D8.1", "created_at": now_iso(),
                "checks": [{"id": identifier, "status": "pass" if passed else "fail", "evidence": evidence} for identifier, passed, evidence in checks],
                "findings": findings, "critical_open": critical_open,
            }
            limitations = [
                "The 28×28 benchmark resolution limits spatial and clinical interpretation.",
                "Class imbalance makes aggregate accuracy insufficient; per-class and macro results control interpretation.",
                "Difference-hash matches are review candidates, not confirmed duplicate identities.",
                "The curated explanation describes a benchmark class and is not individualized medical advice.",
            ]
            decision = "ACCEPTED" if critical_open == 0 else "REJECTED"
            acceptance = {
                "deliverable": "D8.2", "created_at": now_iso(), "decision": decision,
                "critical_findings_open": critical_open, "acceptance_scope": "integrated research software demonstrator",
                "readiness_claim": "TRL 7 target evidence assembled; this is not clinical validation or deployment approval.",
                "accepted_limitations": limitations, "review_checklist": deliverable_path("D8.1").relative_to(ROOT).as_posix(),
            }
            return checklist, acceptance
        checklist, acceptance = self.call("review_evidence", review)
        write_json(deliverable_path("D8.1"), checklist)
        write_json(deliverable_path("D8.2"), acceptance)
        return acceptance


def write_static_deliverables(config: dict[str, Any]) -> None:
    for number in WP_NAMES:
        wp_dir(number)
    review_lines = [
        "# D1.1 Research Review", "", "The project references support five design controls:", "",
        "1. **Dataset integrity:** MedMNIST standardization [1] is paired with exact/perceptual duplicate and split-overlap checks prompted by [2] and dataset documentation from [9].",
        "2. **Balanced evaluation:** aggregate accuracy is accompanied by macro, per-class, uncertainty, and calibration evidence, consistent with the evaluation lessons in [5], [6], and [8].",
        "3. **Bounded explainability:** input-gradient attribution is treated as coarse 28×28 evidence; explanation value and limitations are motivated by [4], [5], [12], and [13].",
        "4. **Constrained agents:** fixed tools, typed handoffs, immutable predictions, and action logs address orchestration and sequential tool risks discussed in [10], [11], [14], [15], and [16].",
        "5. **Research-only scope:** the prototype does not claim the scale, modalities, or clinical validation of [3], [7], or [8].", "", "## References", "",
    ]
    review_lines.extend(f"[{index}] {authors}. “{title}.” {venue} ({year}). {url}" for index, authors, title, venue, year, url, _ in REFERENCES)
    deliverable_path("D1.1").write_text("\n".join(review_lines) + "\n", encoding="utf-8")

    with deliverable_path("D1.2").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Requirement ID", "Area", "Mandatory requirement", "Owner", "Verification method", "Linked deliverable", "Acceptance item", "Definition status"])
        for row in REQUIREMENTS:
            writer.writerow([*row, "Approved"])

    write_json(deliverable_path("D1.3"), {
        "deliverable": "D1.3", "created_at": now_iso(), "status": "approved", "dataset_policy": config["data"],
        "model_scope": {"families": config["training"]["models"], "adaptation": "Compact 28×28 implementations preserve family building blocks while keeping CPU feasibility."},
        "agent_interfaces": {name: {"task": spec[0], "allowed_tools": spec[1], "boundary": spec[2]} for name, spec in AGENT_SPECS.items()},
        "metric_policy": config["evaluation"], "test_rule": "The final test runs after D4.3 is saved and checksummed.",
        "runtime": {"language": "Python 3.12", "ML": "PyTorch", "external_api_required": False},
    })
    write_json(deliverable_path("D3.1"), {
        "deliverable": "D3.1", "status": "approved", "models": config["training"]["models"],
        "shared_controls": {key: config["training"][key] for key in ("seed", "batch_size", "default_max_epochs", "default_target_accuracy", "learning_rate", "weight_decay", "loss", "augmentation")},
        "selection": {"split": "official validation", "primary": config["evaluation"]["primary_metric"], "tie_breakers": config["evaluation"]["tie_breakers"]},
        "architecture_note": "Family-specific compact versions are sized for 28×28 inputs and CPU execution; they are not ImageNet-size parameter replicas.",
    })
    write_json(deliverable_path("D3.2"), {"deliverable": "D3.2", "status": "approved", "agent": "Agent1", "specification": AGENT_SPECS["Agent1"], "input_schema": ["training images/labels", "validation images/labels", "model names", "training settings"], "output_schema": ["run ID", "history", "checkpoint", "checksum", "status/failure reason"], "test_access": "none"})
    write_json(deliverable_path("D4.1"), {"deliverable": "D4.1", "status": "approved_before_selection", **config["evaluation"], "metric_definitions": {"macro_f1": "unweighted mean of seven class F1 values", "balanced_accuracy": "unweighted mean class recall", "macro_auroc": "unweighted mean one-vs-rest rank AUROC", "ece": "10-bin confidence/accuracy gap", "brier": "mean seven-class squared probability error", "nll": "mean negative log probability of true class"}})
    write_json(deliverable_path("D5.1"), {"deliverable": "D5.1", "status": "approved", "approved_source_types": ["peer-reviewed project references", "official MedMNIST documentation", "DermNet clinical topic pages"], "sources": config["sources"], "required_citation_fields": ["title", "url"], "unsupported_claim_behavior": "omit", "separation_rule": "Retrieved context cannot alter predicted class, probabilities, confidence, or uncertainty.", "output_fields": ["trace_id", "predicted_class_id", "fact", "citations", "classifier_mutation"]})
    write_json(deliverable_path("D6.1"), {"deliverable": "D6.1", "status": "approved", "agent": "Agent2", "specification": AGENT_SPECS["Agent2"], "input_schema": ["model manifest", "test configuration", "dataset vault"], "output_schema": ["metrics", "duplicate sensitivity", "robustness", "failure cases"], "mutation_permissions": []})
    write_json(deliverable_path("D6.2"), {"deliverable": "D6.2", "status": "approved_before_test", "test_set": "official DermaMNIST test split", "mandatory_metrics": ["accuracy", "macro_f1", "balanced_accuracy", "macro_auroc", "per-class precision/recall/F1/AUROC", "confusion matrix", "ECE", "Brier score", "NLL"], "robustness": config["evaluation"]["robustness_checks"], "failure_rules": ["checksum mismatch blocks testing", "no post-test tuning", "critical evidence gap blocks acceptance"], "output_format": "versioned JSON plus PNG plots"})
    write_json(deliverable_path("D7.1"), {"deliverable": "D7.1", "status": "approved", "agent": "Agent3", "specification": AGENT_SPECS["Agent3"], "input_schema": ["prediction", "XAI record", "cited class information"], "output_schema": ["prediction", "probabilities", "uncertainty", "attribution", "citations", "trace ID", "research-only notice"], "mutation_check": "probability mapping equality asserted after composition"})


def write_preprocessing_deliverable(bundle: TrainingBundle, config: dict[str, Any]) -> None:
    mean, std = training_statistics(bundle.train_images)
    write_json(deliverable_path("D2.3"), {
        "deliverable": "D2.3", "created_at": now_iso(), "preprocessing_id": "preprocess-v1",
        "input": "28x28 RGB uint8", "scaling": "divide by 255", "normalization_source": "training split only",
        "channel_mean": mean, "channel_std": std, "training_augmentation": config["training"]["augmentation"],
        "validation_augmentation": [], "test_augmentation": [], "loader_batch_size": config["training"]["batch_size"],
        "seed": config["training"]["seed"], "train_indices_sha256": hashlib.sha256(bundle.train_indices.tobytes()).hexdigest(),
        "validation_indices_sha256": hashlib.sha256(bundle.val_indices.tobytes()).hexdigest(), "status": "reproducible",
    })


def create_overview(config: dict[str, Any]) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Install requirements.txt before creating the workbook") from exc

    output = PROJECT_ROOT / "Project_Overview.xlsx"
    workbook = Workbook()
    navy, pale, white = "17365D", "D9EAF7", "FFFFFF"

    def format_sheet(sheet: Any, widths: dict[int, int]) -> None:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color=white, bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index, width in widths.items():
            sheet.column_dimensions[get_column_letter(index)].width = width

    sheet = workbook.active
    sheet.title = "Overview"
    sheet.append(["Section", "Item", "Description"])
    overview_rows = [
        ("Project", "Purpose", "Build and test a multi-agent system for DermaMNIST classification, evaluation, enrichment, explanation, and review."),
        ("Project", "Dataset", "DermaMNIST v2: 10,015 RGB dermoscopic images, 28×28 pixels, seven classes."),
        ("Project", "Task", "Supervised seven-class image classification."),
        ("Project", "Output", "Predicted class, seven probabilities, uncertainty flag, attribution image, cited class information, and trace ID."),
        ("Scope", "Included", "Dataset audit, preprocessing, three classifiers, three agents, metrics, calibration, attribution, citations, final test, and review."),
        ("Scope", "Excluded", "Clinical diagnosis, treatment advice, hospital deployment, private patient data, and changes based on final test results."),
        ("Method", "Development", "Eight work packages. Each work package produces named deliverables and ends with a completion gate."),
        ("Method", "Readiness", "Integrated research software demonstrator. Research and education only."),
    ]
    for row in overview_rows:
        sheet.append(row)
    format_sheet(sheet, {1: 18, 2: 24, 3: 100})

    sheet = workbook.create_sheet("Objectives_RQs")
    sheet.append(["Type", "ID", "Statement", "Expected result"])
    objective_rows = [
        ("Objective", "O1", "Verify the data basis.", "Dataset registry, duplicate checks, class distribution, checksums, and preprocessing settings."),
        ("Objective", "O2", "Build the classifier.", "Three model configurations compared with the same training and evaluation settings."),
        ("Objective", "O3", "Build the multi-agent workflow.", "Agent1, Agent2, and Agent3 have defined inputs, outputs, tools, and responsibilities."),
        ("Objective", "O4", "Verify prediction and explanation.", "Metrics, uncertainty, attribution, citations, and traceable outputs."),
        ("Objective", "O5", "Run and review the prototype.", "Complete run, review checklist, and acceptance report."),
        ("Research question", "RQ1", "Can the agents run reproducible experiments with fixed responsibilities and complete action records?", "Agent records and reproducibility checks."),
        ("Research question", "RQ2", "Which classifier gives the best class-balanced result?", "Validation comparison across the three classifier families."),
        ("Research question", "RQ3", "Can the system add cited class information and visual attribution without changing the classifier result?", "Probability and trace checks in the final report."),
    ]
    for row in objective_rows:
        sheet.append(row)
    format_sheet(sheet, {1: 20, 2: 12, 3: 85, 4: 70})

    sheet = workbook.create_sheet("Work_Packages")
    sheet.append(["WP", "Name", "Requires", "Main work", "Parallel work", "Completion gate"])
    for wp, requires, main, parallel, gate in WORK_PACKAGES:
        sheet.append([wp, WP_NAMES[int(wp[-2:])], requires, main, parallel, gate])
    format_sheet(sheet, {1: 10, 2: 30, 3: 38, 4: 55, 5: 28, 6: 35})

    sheet = workbook.create_sheet("Deliverables")
    sheet.append(["Activity", "Activity name", "Responsible", "Deliverable", "Check"])
    for row in ACTIVITIES:
        sheet.append(row)
    format_sheet(sheet, {1: 12, 2: 34, 3: 25, 4: 14, 5: 85})

    sheet = workbook.create_sheet("Requirements")
    sheet.append(["ID", "Area", "Requirement", "Owner", "Verification", "Deliverable", "Acceptance"])
    for row in REQUIREMENTS:
        sheet.append(row)
    format_sheet(sheet, {1: 18, 2: 18, 3: 72, 4: 24, 5: 48, 6: 28, 7: 16})

    agent_io = {
        "Agent1": ("Training and validation data, model list, stop settings", "Experiment records and model checkpoints"),
        "Agent2": ("Experiment records, validation data, selected model, and test data", "Model selection, metrics, robustness results, and failure cases"),
        "Agent3": ("Prediction, attribution, class sources, and work-package records", "Cited output, review checklist, and acceptance report"),
    }
    sheet = workbook.create_sheet("Agents")
    sheet.append(["Agent", "Task", "Inputs", "Outputs", "Allowed tools", "Limit"])
    for name, (task, tools, boundary) in AGENT_SPECS.items():
        inputs, outputs = agent_io[name]
        sheet.append([name, task, inputs, outputs, ", ".join(tools), boundary])
    format_sheet(sheet, {1: 22, 2: 48, 3: 50, 4: 48, 5: 45, 6: 48})

    sheet = workbook.create_sheet("Technical")
    sheet.append(["Area", "Definition"])
    technical_rows = [
        ("Data", "Official train, validation, and test splits; sample IDs, labels, image shape, class counts, and checksums recorded."),
        ("Data checks", "Exact image hashes, 64-bit difference hashes, label conflicts, and cross-split matches."),
        ("Preprocessing", "RGB values scaled to 0–1, training-channel normalization, and training augmentation."),
        ("Models", "Compact ResNet-18, EfficientNet-B0, and ConvNeXt-Tiny for 28×28 input."),
        ("Training", "Class-weighted cross-entropy, AdamW, fixed seed, early stopping, and saved run history."),
        ("Model selection", "Highest validation macro-F1; balanced accuracy, macro-AUROC, and accuracy break ties."),
        ("Final test", "Runs after model selection with fixed preprocessing, calibration, uncertainty threshold, and metrics."),
        ("Metrics", "Accuracy, macro-F1, balanced accuracy, macro-AUROC, per-class precision/recall/F1, confusion matrix, ECE, Brier score, and NLL."),
        ("Robustness", "Horizontal flip and ±10% brightness checks."),
        ("Attribution", "Absolute input-gradient map from the selected classifier."),
        ("Enrichment", "Class information from the approved source list. Each fact contains a title and URL."),
        ("Output", "JSON prediction report and PNG attribution image."),
        ("Launcher 1", "Run the complete process. Missing DermaMNIST data are downloaded automatically."),
        ("Launcher 2", "Check and test the saved model with one input image."),
        ("Launcher 3", "Open the visual platform."),
        ("Visual platform", "Choose Test trained model or Run full process, upload an image, and view each stage."),
    ]
    for row in technical_rows:
        sheet.append(row)
    format_sheet(sheet, {1: 24, 2: 115})

    sheet = workbook.create_sheet("Risks")
    sheet.append(["ID", "Risk", "Impact", "Trigger", "Mitigation", "Owner"])
    for row in RISKS:
        sheet.append(row)
    format_sheet(sheet, {1: 8, 2: 40, 3: 12, 4: 45, 5: 62, 6: 24})
    for row in sheet.iter_rows(min_row=2):
        if row[2].value in {"Critical", "High"}:
            row[2].fill = PatternFill("solid", fgColor=pale)

    sheet = workbook.create_sheet("Acceptance")
    sheet.append(["ID", "Item", "Deliverables", "Pass condition"])
    for row in ACCEPTANCE:
        sheet.append(row)
    format_sheet(sheet, {1: 8, 2: 38, 3: 18, 4: 90})

    sheet = workbook.create_sheet("References")
    sheet.append(["ID", "Authors", "Title", "Venue", "Year", "DOI or URL", "Project use"])
    for row in REFERENCES:
        sheet.append(row)
    format_sheet(sheet, {1: 8, 2: 32, 3: 72, 4: 34, 5: 10, 6: 55, 7: 42})

    workbook.save(output)
    return output

def create_and_execute_audit_notebook(config: dict[str, Any]) -> Path:
    try:
        import nbformat
        from nbclient import NotebookClient
    except ImportError as exc:
        raise RuntimeError("Install requirements.txt before running the audit notebook") from exc
    if sys.platform == "win32":
        import asyncio
        asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())
    registry, audit = audit_dataset(config, write_outputs=False)
    notebook = nbformat.v4.new_notebook()
    notebook["metadata"]["kernelspec"] = {"display_name": "Python 3", "language": "python", "name": "python3"}
    notebook["cells"] = [
        nbformat.v4.new_markdown_cell(
            f"## tl;dr\n\nThe official file passed MD5, schema, shape, label, and count checks for **{registry['sample_count']:,} images**. "
            f"The audit found **{audit['checks']['cross_split_exact_leakage']['cluster_count']} exact cross-split clusters** and applies parallel leakage-aware test reporting."
        ),
        nbformat.v4.new_markdown_cell("## Context & Methods\n\nThis WP2 notebook reruns the saved audit at one-image-per-record grain.\n\n### Key Assumptions\n\nDifference-hash equality identifies review candidates, not confirmed duplicates."),
        nbformat.v4.new_code_cell("from pathlib import Path\nimport json\nimport derma_agent as project\nconfig = project.load_config()\nprint(config['data']['dataset_id'], project.DATA_PATH)"),
        nbformat.v4.new_markdown_cell("## Data\n\nLoad the official checksummed NPZ and show its fixed partition shapes."),
        nbformat.v4.new_code_cell("arrays = project.load_npz(project.download_dataset(config))\nfor split in ('train', 'val', 'test'):\n    print(split, arrays[f'{split}_images'].shape, arrays[f'{split}_labels'].shape)"),
        nbformat.v4.new_markdown_cell("## Results\n\nRecompute the audit and reconcile its highest-impact counts to saved D2.1/D2.2 evidence."),
        nbformat.v4.new_code_cell(
            "registry, audit = project.audit_dataset(config, write_outputs=False)\n"
            "saved_registry = project.read_json(project.deliverable_path('D2.1'))\n"
            "saved_audit = project.read_json(project.deliverable_path('D2.2'))\n"
            "assert registry['sample_count'] == saved_registry['sample_count'] == 10015\n"
            "for check in ('exact_duplicates', 'cross_split_exact_leakage', 'perceptual_hash_candidates'):\n"
            "    assert audit['checks'][check]['cluster_count'] == saved_audit['checks'][check]['cluster_count']\n"
            "print(json.dumps({k: {'clusters': audit['checks'][k]['cluster_count'], 'records': audit['checks'][k]['affected_records']} for k in ('exact_duplicates', 'cross_split_exact_leakage', 'perceptual_hash_candidates')}, indent=2))"
        ),
        nbformat.v4.new_markdown_cell("## Takeaways\n\n- The official benchmark split remains unchanged for comparability.\n- Test images with exact byte matches in train/validation are excluded in the parallel leakage-aware report.\n- Perceptual-hash candidates remain visible for review and are not silently removed.\n- Training, calibration, uncertainty selection, and model choice never use test labels."),
    ]
    path = wp_dir(2) / "D2.2_reproducible_audit.ipynb"
    client = NotebookClient(notebook, timeout=300, kernel_name="python3", resources={"metadata": {"path": str(ROOT)}})
    executed = client.execute()
    nbformat.write(executed, path)
    return path


def verify_project(config: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    missing = [item for item in DELIVERABLE_FILES if not deliverable_path(item).exists()]
    if missing: errors.append(f"Missing deliverables: {', '.join(missing)}")
    if not DATA_PATH.exists() or data_md5(DATA_PATH) != config["data"]["md5"]: errors.append("Dataset MD5 is missing or invalid")
    if not (PROJECT_ROOT / "Project_Overview.xlsx").exists(): errors.append("Project_Overview.xlsx is missing")
    if errors: return errors
    selection = read_json(deliverable_path("D4.2")); manifest = read_json(deliverable_path("D4.3")); test = read_json(deliverable_path("D6.3")); case = read_json(deliverable_path("D7.2")); prototype = read_json(deliverable_path("D7.3")); checklist = read_json(deliverable_path("D8.1")); acceptance = read_json(deliverable_path("D8.2"))
    if file_hash(saved_model_path(manifest)) != manifest["model_sha256"]: errors.append("Frozen model checksum mismatch")
    expected_selection = max(selection["comparisons"], key=lambda item: selection_key(item, config))
    if selection["selected_model"] != expected_selection["model_name"]: errors.append("Saved selection does not follow the fixed metric order")
    if not selection.get("competitive_on_primary_metric", False): errors.append("Selected validation result does not exceed the project baseline on the primary metric")
    for name in ("accuracy", "macro_f1", "balanced_accuracy", "macro_auroc", "expected_calibration_error"):
        if not 0 <= test["metrics"][name] <= 1: errors.append(f"Metric outside [0,1]: {name}")
    if len(test["metrics"]["confusion_matrix"]) != 7 or any(len(row) != 7 for row in test["metrics"]["confusion_matrix"]): errors.append("Confusion matrix is not 7x7")
    if sum(map(sum, test["metrics"]["confusion_matrix"])) != test["test_sample_count"]: errors.append("Confusion matrix total does not equal tested population")
    if sum(item["support"] for item in test["metrics"]["per_class"]) != test["test_sample_count"]: errors.append("Per-class support does not equal tested population")
    if test["test_sample_count"] != config["data"]["expected_splits"]["test"]: errors.append("Test report does not cover the complete official test split")
    probabilities = list(case["model_evidence"]["probabilities"].values())
    if not math.isclose(sum(probabilities), 1.0, rel_tol=0, abs_tol=1e-6): errors.append("Case probabilities do not sum to one")
    if case["trace_id"] != case["xai_evidence"]["trace_id"]: errors.append("Case/XAI trace mismatch")
    for name, path in prototype["outputs"].items():
        if not (ROOT / path).exists() or file_hash(ROOT / path) != prototype["output_sha256"][name]: errors.append(f"Prototype output checksum mismatch: {name}")
    events = load_events()
    invalid_events = [event for event in events if event["status"] == "completed" and event["tool"] not in AGENT_SPECS.get(event["agent"], (None, [], None))[1]]
    if invalid_events: errors.append("Completed action log contains a tool outside its agent allow list")
    if not all(item["status"] == "pass" for item in checklist["checks"]): errors.append("At least one final acceptance check failed")
    if acceptance["critical_findings_open"] != 0: errors.append("Critical findings remain open")
    if not (wp_dir(2) / "D2.2_reproducible_audit.ipynb").exists(): errors.append("Executed WP2 audit notebook is missing")

    arrays = load_npz(DATA_PATH)
    for split, expected in config["data"]["expected_splits"].items():
        if len(arrays[f"{split}_images"]) != expected: errors.append(f"Raw {split} count changed")
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    record = {"model_name": manifest["model_name"], "model_file": manifest["model_file"], "model_sha256": manifest["model_sha256"]}
    model = load_checkpoint_model(record, device)
    logits, labels = predict_logits(model, arrays["test_images"], arrays["test_labels"].reshape(-1), np.array(manifest["normalization_mean"]), np.array(manifest["normalization_std"]), config["training"]["batch_size"], device)
    recalculated = classification_metrics(labels, softmax_numpy(logits, manifest["temperature"]), config["evaluation"]["calibration_bins"])
    for name in ("accuracy", "macro_f1", "balanced_accuracy", "macro_auroc", "expected_calibration_error", "brier_score", "negative_log_likelihood"):
        if not math.isclose(recalculated[name], test["metrics"][name], rel_tol=0, abs_tol=1e-10): errors.append(f"Test metric recomputation differs: {name}")
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(PROJECT_ROOT / "Project_Overview.xlsx", data_only=False, read_only=False)
        expected_sheets = {"Overview", "Objectives_RQs", "Work_Packages", "Deliverables", "Requirements", "Agents", "Technical", "Risks", "Acceptance", "References"}
        if set(workbook.sheetnames) != expected_sheets: errors.append("Workbook sheet set is incomplete")
        forbidden_headers = {"status", "progress", "evidence file", "action log", "current state"}
        headers = {str(cell.value).lower() for sheet in workbook for cell in sheet[1] if cell.value is not None}
        if headers & forbidden_headers: errors.append("Workbook contains progress-tracking fields")
        if any(cell.hyperlink for sheet in workbook for row in sheet.iter_rows() for cell in row): errors.append("Workbook contains file or web hyperlinks")
        if any(isinstance(cell.value, str) and cell.value.startswith("=") for sheet in workbook for row in sheet.iter_rows() for cell in row): errors.append("Workbook contains formulas")
        workbook.close()
    except Exception as exc:
        errors.append(f"Workbook validation failed: {exc}")
    return errors


def resolve_training_limits(
    config: dict[str, Any],
    target_accuracy: float | None,
    max_epochs: int | None,
    confirmed: bool,
) -> tuple[float, int]:
    default_target = float(config["training"]["default_target_accuracy"])
    default_epochs = int(config["training"]["default_max_epochs"])
    if target_accuracy is None:
        entered = input(f"Stop when validation accuracy reaches (0-1) [{default_target:.2f}]: ").strip()
        target_accuracy = default_target if not entered else float(entered)
    if max_epochs is None:
        entered = input(f"Maximum training epochs [{default_epochs}]: ").strip()
        max_epochs = default_epochs if not entered else int(entered)
    if not 0 < target_accuracy <= 1:
        raise ValueError("Target accuracy must be greater than 0 and no more than 1")
    if not 1 <= max_epochs <= 100:
        raise ValueError("Maximum epochs must be from 1 to 100")
    if not confirmed:
        answer = input(f"Start training with target {target_accuracy:.2f} and at most {max_epochs} epochs? [y/N]: ").strip().lower()
        if answer not in {"y", "yes"}:
            raise ValueError("Training cancelled")
    return target_accuracy, max_epochs


def export_full_process_results(
    records: list[dict[str, Any]],
    acceptance: dict[str, Any],
    prediction: dict[str, Any],
    output_dir: Path = FULL_PROCESS_OUTPUT,
) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    exports = {
        "training_results.json": deliverable_path("D3.3"),
        "model_selection.json": deliverable_path("D4.2"),
        "final_test_results.json": deliverable_path("D6.3"),
        "review.json": deliverable_path("D8.2"),
        "confusion_matrix.png": wp_dir(6) / "confusion_matrix.png",
        "calibration_plot.png": wp_dir(6) / "calibration_plot.png",
    }
    for name, source in exports.items():
        shutil.copy2(source, output_dir / name)
    test = read_json(deliverable_path("D6.3"))
    selection = read_json(deliverable_path("D4.2"))
    summary = {
        "completed_at": now_iso(),
        "selected_model": selection["selected_model"],
        "training_stops": {record["model_name"]: record["stop_reason"] for record in records},
        "test_accuracy": test["metrics"]["accuracy"],
        "test_macro_f1": test["metrics"]["macro_f1"],
        "review_decision": acceptance["decision"],
        "input_prediction": prediction["prediction"],
        "files": ["prediction.json", "attribution.png", *exports],
    }
    write_json(output_dir / "run_summary.json", summary)
    return summary


def run_full_process(
    image_path: Path,
    target_accuracy: float,
    max_epochs: int,
    output_dir: Path = FULL_PROCESS_OUTPUT,
    progress: Callable[[int, str], None] | None = None,
) -> dict[str, Any]:
    validate_input_image(image_path)
    config = load_config()
    set_seed(config["training"]["seed"])
    REGISTRY_PATH.unlink(missing_ok=True)
    report_progress(progress, 3, "Creating work-package specifications")
    print("[1/8] Creating work-package specifications")
    write_static_deliverables(config)
    report_progress(progress, 8, "Checking and auditing DermaMNIST")
    print("[2/8] Auditing the DermaMNIST dataset")
    _, audit = audit_dataset(config)
    vault = DatasetVault(config)
    bundle = vault.training_bundle()
    write_preprocessing_deliverable(bundle, config)
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    report_progress(progress, 18, "Agent1 is starting model training")
    print("[3/8] Agent1 is training the three models")
    records = TrainingAgent().run(bundle, config, target_accuracy, max_epochs, device, progress)
    report_progress(progress, 58, "Agent2 is comparing the trained models")
    print("[4/8] Agent2 is selecting the model")
    manifest = EvaluationAgent().run(records, bundle, config, device)
    report_progress(progress, 68, "Agent2 is running the final test")
    print("[5/8] Agent2 is running the final test")
    TestingAgent().run(vault, manifest, audit, config, device)
    report_progress(progress, 82, "Agent3 is processing the input image")
    print("[6/8] Agent3 is processing the input image")
    prediction = run_user_image(image_path, manifest, config, device, output_dir)
    report_progress(progress, 92, "Agent3 is reviewing the project records")
    print("[7/8] Agent3 is reviewing the project records")
    acceptance = ReviewAgent().run(config)
    report_progress(progress, 97, "Exporting the complete result")
    print("[8/8] Exporting results")
    create_overview(config)
    summary = export_full_process_results(records, acceptance, prediction, output_dir)
    report_progress(progress, 100, "Complete process finished")
    return {"summary": summary, "acceptance": acceptance, "prediction": prediction}


def run_saved_model(image_path: Path) -> dict[str, Any]:
    validate_input_image(image_path)
    config = load_config()
    manifest_path = deliverable_path("D4.3")
    if not manifest_path.exists():
        raise FileNotFoundError("No trained model was found. Run the full process first.")
    result = run_user_image(
        image_path,
        read_json(manifest_path),
        config,
        torch.device("cuda" if torch.cuda.is_available() else "cpu"),
        FINAL_OUTPUT,
    )
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="DermaAgent")
    subparsers = parser.add_subparsers(dest="command", required=True)
    test_parser = subparsers.add_parser("test", help="Test the saved model with one image")
    test_parser.add_argument("image", nargs="?", default=str(PROJECT_ROOT / "input" / "sample_derma.png"))
    full_parser = subparsers.add_parser("full", help="Train, select, test, and process one image")
    full_parser.add_argument("image", nargs="?", default=str(PROJECT_ROOT / "input" / "sample_derma.png"))
    full_parser.add_argument("--target-accuracy", type=float)
    full_parser.add_argument("--max-epochs", type=int)
    full_parser.add_argument("--yes", action="store_true", help="Use the supplied stop settings without confirmation")
    args = parser.parse_args(argv)
    try:
        if args.command == "test":
            result = run_saved_model(Path(args.image))
            print(json.dumps({key: result[key] for key in ("prediction", "confidence", "uncertain", "trace_id")}, indent=2))
            print(f"Output: {FINAL_OUTPUT}")
            return 0
        target, maximum = resolve_training_limits(load_config(), args.target_accuracy, args.max_epochs, args.yes)
        result = run_full_process(Path(args.image), target, maximum)
        print(json.dumps(result["summary"], indent=2))
        print(f"Output: {FULL_PROCESS_OUTPUT}")
        return 0 if result["acceptance"]["critical_findings_open"] == 0 else 2
    except (EOFError, FileNotFoundError, ValueError, IndexError, OSError, RuntimeError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
